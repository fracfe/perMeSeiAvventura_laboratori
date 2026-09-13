from functools import wraps

from flask import Flask, render_template, redirect, request, url_for, flash, send_file, session, abort
from flask_login import UserMixin, login_user, LoginManager, login_required, logout_user, current_user
from werkzeug.security import generate_password_hash, check_password_hash
from werkzeug.exceptions import HTTPException
from werkzeug.middleware.proxy_fix import ProxyFix
from flask_sqlalchemy import SQLAlchemy
from flask_migrate import Migrate
from sqlalchemy import func, or_
from sqlalchemy.exc import IntegrityError, SQLAlchemyError
from sqlalchemy.orm import aliased
from openpyxl import Workbook
from datetime import datetime
from zoneinfo import ZoneInfo
from suddivisione_gruppi_service import (
    NOMI_GRUPPI_DOMENICA,
    NUMERO_GRUPPI_DOMENICA,
    calcola_gruppi_domenica,
    calcola_sottogruppi_ab,
    ErroreSuddivisione,
    carica_excel_suddivisione,
    distribuzione_dimensioni,
    estrai_partecipanti_suddivisione,
    genera_excel_suddivisione,
    valida_numero_gruppi,
)
from comunicazioni_service import crea_workbook_comunicazioni, descrivi_domenica_comunicazione

import json
import io
import os
import re

drivers = {
    "sqlite": "sqlite:///",
    "mariadb": "mysql+pymysql://",
}

# Inizializza app e servizi
app = Flask(__name__)
db_type = os.environ["DB_TYPE"]

if db_type not in drivers:
    app.logger.info("Tipo di database non supportato")
    raise RuntimeError("Tipo di database non supportato")

if db_type == "sqlite":
    uri = f"{drivers[db_type]}{os.environ['DB_NAME']}"
else:
    uri = (
        f"{drivers[db_type]}"
        f"{os.environ['DB_USER']}:"
        f"{os.environ['DB_PASSWORD']}@"
        f"{os.environ['DB_HOST']}:"
        f"{os.environ['DB_PORT']}/"
        f"{os.environ['DB_NAME']}"
    )
app.config["SQLALCHEMY_DATABASE_URI"] =  uri
app.config["SECRET_KEY"] = os.environ['SECRET_KEY']
app.config["MAX_CONTENT_LENGTH"] = 10 * 1024 * 1024
app.config["SESSION_COOKIE_HTTPONLY"] = True
app.config["SESSION_COOKIE_SAMESITE"] = "Lax"
app.config["SESSION_COOKIE_SECURE"] = os.environ.get(
    "SESSION_COOKIE_SECURE",
    "false",
).strip().casefold() in {"1", "true", "yes", "on"}
if db_type == "mariadb":
    app.config["SQLALCHEMY_ENGINE_OPTIONS"] = {"isolation_level": "READ COMMITTED"}
app.wsgi_app = ProxyFix(app.wsgi_app, x_for=1, x_proto=1, x_host=1)
db = SQLAlchemy(app)

login_manager = LoginManager(app)
login_manager.login_view = "index"
login_manager.login_message = u"Sessione scaduta!"


def admin_required(*, api=False):
    def decorator(funzione):
        @wraps(funzione)
        @login_required
        def funzione_protetta(*args, **kwargs):
            if current_user.username != "admin":
                if api:
                    return {"ok": False, "errore": "Accesso non autorizzato."}, 403
                return redirect(url_for("index"))
            return funzione(*args, **kwargs)

        return funzione_protetta

    return decorator

# Classi Database
class User(db.Model, UserMixin):
    __tablename__ = "user"
    id = db.Column(db.Integer, primary_key=True)
    username = db.Column(db.String(255), nullable=False, unique=True)
    password = db.Column(db.String(255), nullable=False)

class Iscrizione(db.Model):
    __tablename__ = "iscrizioni"
    __table_args__ = (
        db.UniqueConstraint("partecipante", name="uq_iscrizioni_partecipante"),
        db.CheckConstraint(
            "scelta_mattino IS NULL OR non_partecipa_mattino = 0",
            name="ck_iscrizioni_scelta_mattino_esclusiva",
        ),
        db.CheckConstraint(
            "scelta_pomeriggio IS NULL OR non_partecipa_pomeriggio = 0",
            name="ck_iscrizioni_scelta_pomeriggio_esclusiva",
        ),
        db.CheckConstraint(
            "sottogruppo_mattino IS NULL OR sottogruppo_mattino IN ('A', 'B')",
            name="ck_iscrizioni_sottogruppo_mattino",
        ),
        db.CheckConstraint(
            "sottogruppo_pomeriggio IS NULL OR sottogruppo_pomeriggio IN ('A', 'B')",
            name="ck_iscrizioni_sottogruppo_pomeriggio",
        ),
    )
    id = db.Column(db.Integer, primary_key=True)
    data = db.Column(db.DateTime, nullable=False)
    partecipante = db.Column(db.Integer, db.ForeignKey("partecipanti.id", name="fk_iscrizioni_partecipanti_id"), nullable=False)
    scelta_mattino = db.Column(db.Integer, db.ForeignKey("laboratori.id", name="fk_iscrizioni_laboratori_mattino_id"), nullable=True)
    scelta_pomeriggio = db.Column(db.Integer, db.ForeignKey("laboratori.id", name="fk_iscrizioni_laboratori_pomeriggio_id"), nullable=True)
    non_partecipa_mattino = db.Column(db.Boolean, nullable=False, default=False)
    non_partecipa_pomeriggio = db.Column(db.Boolean, nullable=False, default=False)
    sottogruppo_mattino = db.Column(db.String(1), nullable=True)
    sottogruppo_pomeriggio = db.Column(db.String(1), nullable=True)

class Partecipante(db.Model):
    __tablename__ = "partecipanti"
    __table_args__ = (
        db.CheckConstraint(
            "gruppo_domenica IS NULL OR gruppo_domenica BETWEEN 1 AND 20",
            name="ck_partecipanti_gruppo_domenica",
        ),
    )
    id = db.Column(db.Integer, primary_key=True)
    nome = db.Column(db.String(255), nullable=False)
    cognome = db.Column(db.String(255), nullable=False)
    gruppo = db.Column(db.String(255), nullable=True)
    zona = db.Column(db.String(255), nullable=True)
    regione = db.Column(db.String(255), nullable=True)
    email = db.Column(db.String(255), nullable=True)
    sesso = db.Column(db.String(50), nullable=True)
    foca = db.Column(db.String(255), nullable=True)
    ruolo = db.Column(db.String(255), nullable=True)
    incarico_altro = db.Column(db.Text, nullable=True)
    deve_iscriversi_sabato = db.Column(
        db.Boolean, nullable=False, default=True, server_default=db.true()
    )
    includi_domenica = db.Column(
        db.Boolean, nullable=False, default=False, server_default=db.false()
    )
    gruppo_domenica = db.Column(db.SmallInteger, nullable=True)

class Laboratorio(db.Model):
    __tablename__ = "laboratori"
    id = db.Column(db.Integer, primary_key=True)
    id_lab = db.Column(db.String(10), nullable=False)
    titolo = db.Column(db.String(255), nullable=False)
    descrizione = db.Column(db.Text, nullable=False)
    posti = db.Column(db.Integer, nullable=False)
    tipologia = db.Column(db.String(10), nullable=False)

class SysOption(db.Model):
    __tablename__ = "system_option"
    key = db.Column(db.String(128), primary_key=True)
    value = db.Column(db.String(128), nullable=False)

migrate = Migrate(app, db)

STATO_ISCRIZIONI_KEY = "stato_iscrizioni"
MESSAGGIO_ISCRIZIONI_KEY = "messaggio_iscrizioni"
ULTIMO_IMPORT_PARTECIPANTI_KEY = "ultimo_import_partecipanti"
ULTIMO_IMPORT_LABORATORI_KEY = "ultimo_import_laboratori"
ULTIMO_RICALCOLO_DOMENICA_KEY = "ultimo_ricalcolo_domenica"
STATI_ISCRIZIONI_VALIDI = ("aperte", "chiuse")
SCELTA_NON_PARTECIPA = "non_partecipa"
TESTO_NON_PARTECIPA = "Non partecipa"
MAX_INTEGER_DATABASE = 2147483647
MAX_TESTO_DATABASE = 65535

def get_stato_iscrizioni():
    opzione = db.session.get(SysOption, STATO_ISCRIZIONI_KEY)
    if opzione is None or opzione.value not in STATI_ISCRIZIONI_VALIDI:
        return "chiuse"
    return opzione.value

def get_messaggio_iscrizioni():
    opzione = db.session.get(SysOption, MESSAGGIO_ISCRIZIONI_KEY)
    if opzione is None:
        return ""
    return opzione.value

def set_sys_option(key, value):
    opzione = db.session.get(SysOption, key)
    if opzione is None:
        db.session.add(SysOption(key=key, value=value))
    else:
        opzione.value = value

def registra_ultimo_import(key):
    data_import = datetime.now(ZoneInfo("Europe/Rome")).isoformat()
    set_sys_option(key, data_import)

def get_ultimo_import(key):
    opzione = db.session.get(SysOption, key)
    if opzione is None:
        return "mai"
    try:
        data_import = datetime.fromisoformat(opzione.value)
    except ValueError:
        return "mai"
    return data_import.strftime("%d/%m/%Y %H:%M")

COLONNE_CSV_PARTECIPANTI = {
    "Codice": "id",
    "Nome": "nome",
    "Cognome": "cognome",
    "Gruppo": "gruppo",
    "Zona": "zona",
    "Regione": "regione",
    "EmailContatto": "email",
    "Sesso": "sesso",
    "FoCa": "foca",
    "Partecipo in qualità di:": "ruolo",
    'Se hai indicato "altro" specifica incarico:': "incarico_altro",
    "Partecipa ai laboratori di sabato come partecipante": "deve_iscriversi_sabato",
    "Partecipa ai laboratori di domenica come partecipante": "includi_domenica",
}
ETICHETTE_CAMPI_PARTECIPANTE = {campo: etichetta for etichetta, campo in COLONNE_CSV_PARTECIPANTI.items()}
ETICHETTE_CAMPI_PARTECIPANTE.update(email="Email", ruolo="Ruolo", incarico_altro="Incarico altro")
CAMPI_MANUALI_OBBLIGATORI = ("nome", "cognome", "gruppo", "zona", "regione", "email", "sesso", "foca")

CAMPI_BOOLEANI_PARTECIPANTE = ("deve_iscriversi_sabato", "includi_domenica")


def normalizza_flag_partecipante(valore):
    if isinstance(valore, bool):
        return valore
    if isinstance(valore, str):
        valore = valore.strip().casefold()
        if valore in {"sì", "si", "s", "yes", "true", "1"}:
            return True
        if valore in {"no", "n", "false", "0"}:
            return False
    raise ValueError("Il flag deve contenere un valore esplicito sì/no valido.")


def valida_import_partecipanti(dati):
    if not isinstance(dati, list):
        return None, "Il payload deve contenere una lista di partecipanti."
    if not dati:
        return None, "Il file deve contenere almeno un partecipante."

    partecipanti_validati = []
    codici_visti = set()
    for indice, riga in enumerate(dati, start=1):
        if not isinstance(riga, dict):
            return None, f"La riga {indice} non contiene un partecipante valido."
        for campo in COLONNE_CSV_PARTECIPANTI.values():
            if campo not in riga:
                return None, f'Manca il campo "{campo}" nella riga {indice}.'

        partecipante_id = riga["id"]
        if (
            isinstance(partecipante_id, bool)
            or not isinstance(partecipante_id, int)
            or partecipante_id <= 0
            or partecipante_id > MAX_INTEGER_DATABASE
        ):
            return None, f"La riga {indice} contiene un codice censimento non valido."
        if partecipante_id in codici_visti:
            return None, f"La riga {indice} contiene il codice censimento duplicato {partecipante_id}."
        codici_visti.add(partecipante_id)
        validata = {"id": partecipante_id}
        for campo in COLONNE_CSV_PARTECIPANTI.values():
            if campo == "id":
                continue
            valore = riga[campo]
            if campo in CAMPI_BOOLEANI_PARTECIPANTE:
                try:
                    validata[campo] = normalizza_flag_partecipante(valore)
                except ValueError:
                    return None, f'Il flag "{campo}" nella riga {indice} deve contenere un valore esplicito sì/no valido.'
                continue
            obbligatorio = campo in ("nome", "cognome")
            if valore is None and not obbligatorio:
                validata[campo] = None
                continue
            if not isinstance(valore, str):
                return None, f'Il campo "{campo}" nella riga {indice} deve essere un testo.'
            valore = valore.strip()
            if not valore and obbligatorio:
                return None, f'Manca il campo "{campo}" nella riga {indice}.'
            if campo == "incarico_altro":
                if len(valore.encode("utf-8")) > MAX_TESTO_DATABASE:
                    return None, f'Il campo "{campo}" nella riga {indice} è troppo lungo.'
            else:
                limite = 50 if campo == "sesso" else 255
                if len(valore) > limite:
                    return None, f'Il campo "{campo}" nella riga {indice} supera {limite} caratteri.'
            validata[campo] = valore or None
        partecipanti_validati.append(validata)
    return partecipanti_validati, None

def valida_laboratorio(riga, riferimento="scheda del laboratorio"):
    if not isinstance(riga, dict):
        return None, f"La {riferimento} non è valida."

    for campo in ("id", "titolo", "descrizione", "posti"):
        if campo not in riga:
            return None, f'Manca il campo "{campo}" nella {riferimento}.'

    id_lab = riga["id"]
    if not isinstance(id_lab, str) or not id_lab.strip():
        return None, f"Manca il codice laboratorio nella {riferimento}."
    id_lab = id_lab.strip()
    if len(id_lab) > 10:
        return None, f"Il codice laboratorio nella {riferimento} supera 10 caratteri."

    titolo = riga["titolo"]
    if not isinstance(titolo, str) or not titolo.strip():
        return None, f"Manca il titolo nella {riferimento}."
    titolo = titolo.strip()
    if len(titolo) > 255:
        return None, f"Il titolo nella {riferimento} supera 255 caratteri."

    descrizione = riga["descrizione"]
    if not isinstance(descrizione, str) or not descrizione.strip():
        return None, f"Manca la descrizione nella {riferimento}."
    descrizione = descrizione.strip()
    if len(descrizione.encode("utf-8")) > MAX_TESTO_DATABASE:
        return None, f"La descrizione nella {riferimento} è troppo lunga."

    posti = riga["posti"]
    if (
        isinstance(posti, bool)
        or not isinstance(posti, int)
        or posti <= 0
        or posti > MAX_INTEGER_DATABASE
    ):
        return None, f"Il laboratorio {id_lab} ha una capienza non valida."

    return {"id": id_lab, "titolo": titolo, "descrizione": descrizione, "posti": posti}, None

def valida_import_laboratori(dati):
    if not isinstance(dati, dict):
        return None, "Il payload dei laboratori non è valido."

    laboratori_validati = {}
    for chiave, fascia in (
        ("lab_mattino", "mattino"),
        ("lab_pomeriggio", "pomeriggio"),
    ):
        if chiave not in dati:
            return None, f'Manca l\'elenco "{chiave}".'

        righe = dati[chiave]
        if not isinstance(righe, list):
            return None, f'L\'elenco "{chiave}" deve essere una lista.'
        if not righe:
            return None, (
                "Il file deve contenere almeno un laboratorio del mattino "
                "e uno del pomeriggio."
            )

        righe_validate = []
        codici_visti = set()
        for indice, riga in enumerate(righe, start=1):
            riferimento = f"riga {indice} dei laboratori del {fascia}"
            validata, errore = valida_laboratorio(riga, riferimento)
            if errore:
                return None, errore
            if validata["id"] in codici_visti:
                return None, f"Codice laboratorio duplicato {validata['id']} nel foglio {fascia}."
            codici_visti.add(validata["id"])
            righe_validate.append(validata)

        laboratori_validati[chiave] = righe_validate

    return laboratori_validati, None

def get_partecipante_corrente(richiedi_conferma=True):
    if not current_user.is_authenticated:
        return None
    if richiedi_conferma and not session.get("identita_confermata", False):
        return None
    user_id = str(current_user.get_id())
    if not user_id.startswith("temp:"):
        return None
    try:
        partecipante_id = int(user_id.removeprefix("temp:"))
    except ValueError:
        return None
    return db.session.get(Partecipante, partecipante_id)

def get_iscrizione_partecipante(partecipante_id):
    return Iscrizione.query.filter_by(partecipante=partecipante_id).first()

def iscrizione_completa(iscrizione):
    return (
        iscrizione is not None
        and scelta_fascia_effettuata(iscrizione, "mattino")
        and scelta_fascia_effettuata(iscrizione, "pomeriggio")
    )

def scelta_fascia_effettuata(iscrizione, tipologia):
    return (
        iscrizione is not None
        and (
            getattr(iscrizione, f"scelta_{tipologia}") is not None
            or getattr(iscrizione, f"non_partecipa_{tipologia}")
        )
    )

def stato_iscrizione(iscrizione, partecipante=None):
    if partecipante is not None and not partecipante.deve_iscriversi_sabato:
        return "non_richiesta"
    if iscrizione_completa(iscrizione):
        return "completo"
    if scelta_fascia_effettuata(iscrizione, "mattino") != scelta_fascia_effettuata(
        iscrizione, "pomeriggio"
    ):
        return "incompleto"
    return "non_iniziato"

def ora_roma():
    return datetime.now(ZoneInfo("Europe/Rome")).replace(tzinfo=None)

class TemporaryUser(UserMixin):
    def __init__(self, id):
        self.id = id
        self.username = id

@app.cli.command("init_db")
def init_db():
    try:
        db.session.add(User(username="admin", password=generate_password_hash("password")))
        print("Utente 'admin' creato con password: 'password'")
        db.session.commit()
        print("Operazione terminata correttamente!")
    except Exception as e:
        print("Qualcosa è andato storto!")
        print(e)

@app.cli.command("reset")
def reset_tool():
    Iscrizione.query.delete()
    Partecipante.query.delete()
    Laboratorio.query.delete()
    db.session.commit()
    print("Reset tool!")

@login_manager.user_loader
def load_user(user_id):
    # Utente temporaneo
    if user_id.startswith("temp:"):
        temp_data = session.get("temp_user")
        if temp_data and temp_data["id"] == user_id:
            return TemporaryUser(
                temp_data["id"]
            )
        return None

    # Utente permanente
    return User.query.get(int(user_id))

@app.route("/")
def index():
    return render_template(
        "index.html",
        stato_iscrizioni=get_stato_iscrizioni(),
        messaggio_iscrizioni=get_messaggio_iscrizioni(),
    )

@app.route("/verifica_iscrizione", methods=["POST"])
def verifica_iscrizione():
    dati = request.get_json(silent=True)
    output = {
        "stato_verifica": False,
        "nome": "",
        "cognome": "",
    }
    if not dati or "codice_socio" not in dati:
        return {**output, "errore": "Inserisci un codice censimento valido."}, 400

    try:
        partecipante_id = int(dati["codice_socio"])
    except (TypeError, ValueError):
        partecipante_id = None

    partecipante = db.session.get(Partecipante, partecipante_id)
    if partecipante is None:
        if str(session.get("_user_id", "")).startswith("temp:"):
            logout_user()
        session.pop("temp_user", None)
        session.pop("identita_confermata", None)
        return {**output, "errore": "Codice censimento non trovato."}, 404

    output["stato_verifica"] = True
    output["nome"] = partecipante.nome
    output["cognome"] = partecipante.cognome

    temp_id = f"temp:{partecipante.id}"
    session["temp_user"] = {
        "id": temp_id
    }
    session["identita_confermata"] = False
    login_user(
        TemporaryUser(
            temp_id
        )
    )
    return output

@app.route("/conferma_identita", methods=["POST"])
def conferma_identita():
    partecipante = get_partecipante_corrente(richiedi_conferma=False)
    if partecipante is None:
        return redirect(url_for("index"))
    session["identita_confermata"] = True
    if not partecipante.deve_iscriversi_sabato:
        return redirect(url_for("riepilogo_iscrizione"))
    return redirect(url_for("percorso_iscrizione"))

@app.route("/iscrizione")
def percorso_iscrizione():
    partecipante = get_partecipante_corrente()
    if partecipante is None:
        return redirect(url_for("index"))
    if not partecipante.deve_iscriversi_sabato:
        return redirect(url_for("riepilogo_iscrizione"))

    iscrizione = get_iscrizione_partecipante(partecipante.id)
    if get_stato_iscrizioni() == "chiuse":
        return redirect(url_for("riepilogo_iscrizione"))
    if not scelta_fascia_effettuata(iscrizione, "mattino"):
        return redirect(url_for("scelta_laboratorio", tipologia="mattino"))
    if not scelta_fascia_effettuata(iscrizione, "pomeriggio"):
        return redirect(url_for("scelta_laboratorio", tipologia="pomeriggio"))
    return redirect(url_for("riepilogo_iscrizione"))

@app.route("/laboratori")
def laboratori():
    return redirect(url_for("percorso_iscrizione"))

@app.route("/laboratori/<tipologia>")
def scelta_laboratorio(tipologia):
    if tipologia not in ("mattino", "pomeriggio"):
        return render_template("errore_generico.html"), 404

    partecipante = get_partecipante_corrente()
    if partecipante is None:
        return redirect(url_for("index"))
    if not partecipante.deve_iscriversi_sabato:
        flash("Iscrizione ai laboratori del sabato: non richiesta", "info")
        return redirect(url_for("riepilogo_iscrizione"))
    if get_stato_iscrizioni() != "aperte":
        return redirect(url_for("riepilogo_iscrizione"))

    iscrizione = get_iscrizione_partecipante(partecipante.id)
    modifica = request.args.get("modifica") == "1"

    if tipologia == "pomeriggio" and (
        not scelta_fascia_effettuata(iscrizione, "mattino")
    ):
        return redirect(url_for("scelta_laboratorio", tipologia="mattino"))

    scelta_corrente = None
    if iscrizione is not None:
        if getattr(iscrizione, f"non_partecipa_{tipologia}"):
            scelta_corrente = SCELTA_NON_PARTECIPA
        else:
            scelta_corrente = getattr(iscrizione, f"scelta_{tipologia}")

    if modifica:
        if not iscrizione_completa(iscrizione) or not scelta_fascia_effettuata(
            iscrizione, tipologia
        ):
            return redirect(url_for("percorso_iscrizione"))
    elif scelta_corrente is not None:
        return redirect(url_for("percorso_iscrizione"))

    return render_template(
        "laboratori.html",
        tipologia=tipologia,
        scelta_corrente=scelta_corrente,
        scelta_non_partecipa=SCELTA_NON_PARTECIPA,
        modifica=modifica,
    )

@app.route("/lista_laboratori/<tipologia>")
def lista_laboratori(tipologia):
    if tipologia not in ("mattino", "pomeriggio"):
        return {"ok": False, "errore": "Tipologia laboratorio non valida."}, 400

    partecipante = get_partecipante_corrente()
    if partecipante is None:
        return {"ok": False, "errore": "Sessione scaduta."}, 401
    if not partecipante.deve_iscriversi_sabato:
        return {"ok": False, "errore": "Iscrizione ai laboratori del sabato: non richiesta"}, 403
    if get_stato_iscrizioni() != "aperte":
        return {"ok": False, "errore": "Le iscrizioni sono chiuse."}, 403

    iscrizione = get_iscrizione_partecipante(partecipante.id)
    scelta_corrente = (
        SCELTA_NON_PARTECIPA
        if iscrizione and getattr(iscrizione, f"non_partecipa_{tipologia}")
        else getattr(iscrizione, f"scelta_{tipologia}") if iscrizione else None
    )
    colonna_scelta = getattr(Iscrizione, f"scelta_{tipologia}")
    conteggi = dict(
        db.session.query(colonna_scelta, func.count(Iscrizione.id))
        .filter(colonna_scelta.is_not(None))
        .group_by(colonna_scelta)
        .all()
    )

    laboratori_output = [
        {
            "id": SCELTA_NON_PARTECIPA,
            "id_lab": "",
            "titolo": "Non partecipo a nessun laboratorio",
            "descrizione": "Seleziona questa opzione se non parteciperai in questa fascia.",
            "posti": None,
            "posti_disponibili": None,
            "posseduto": scelta_corrente == SCELTA_NON_PARTECIPA,
            "selezionabile": True,
            "speciale": True,
        }
    ]
    for laboratorio in Laboratorio.query.filter_by(tipologia=tipologia).order_by(
        Laboratorio.id_lab
    ):
        occupati = conteggi.get(laboratorio.id, 0)
        posseduto = scelta_corrente == laboratorio.id
        posti_disponibili = max(laboratorio.posti - occupati, 0)
        laboratori_output.append(
            {
                "id": laboratorio.id,
                "id_lab": laboratorio.id_lab,
                "titolo": laboratorio.titolo,
                "descrizione": laboratorio.descrizione,
                "posti": laboratorio.posti,
                "posti_disponibili": posti_disponibili,
                "posseduto": posseduto,
                "selezionabile": posseduto or posti_disponibili > 0,
                "speciale": False,
            }
        )
    return {"ok": True, "laboratori": laboratori_output}

@app.route("/laboratori/<tipologia>/salva", methods=["POST"])
def salva_laboratorio(tipologia):
    if tipologia not in ("mattino", "pomeriggio"):
        return {"ok": False, "errore": "Tipologia laboratorio non valida."}, 400
    if get_stato_iscrizioni() != "aperte":
        return {"ok": False, "errore": "Le iscrizioni sono chiuse."}, 403

    partecipante_sessione = get_partecipante_corrente()
    if partecipante_sessione is None:
        return {"ok": False, "errore": "Sessione scaduta."}, 401
    if not partecipante_sessione.deve_iscriversi_sabato:
        return {"ok": False, "errore": "Iscrizione ai laboratori del sabato: non richiesta"}, 403

    dati = request.get_json(silent=True)
    if not dati:
        return {"ok": False, "errore": "Seleziona un laboratorio valido."}, 400
    non_partecipa = dati.get("non_partecipa") is True
    laboratorio_id = None
    if non_partecipa:
        if "laboratorio_id" in dati:
            return {"ok": False, "errore": "Scelta non valida."}, 400
    else:
        if "laboratorio_id" not in dati:
            return {"ok": False, "errore": "Seleziona un laboratorio valido."}, 400
        try:
            laboratorio_id = int(dati["laboratorio_id"])
        except (TypeError, ValueError):
            return {"ok": False, "errore": "Laboratorio non valido."}, 400

    try:
        partecipante = db.session.execute(
            db.select(Partecipante)
            .where(Partecipante.id == partecipante_sessione.id)
            .with_for_update()
            .execution_options(populate_existing=True)
        ).scalar_one_or_none()
        if partecipante is None:
            db.session.rollback()
            return {"ok": False, "errore": "Partecipante non valido."}, 400
        if not partecipante.deve_iscriversi_sabato:
            db.session.rollback()
            return {"ok": False, "errore": "Iscrizione ai laboratori del sabato: non richiesta"}, 403

        iscrizione = get_iscrizione_partecipante(partecipante.id)
        if tipologia == "pomeriggio" and (
            not scelta_fascia_effettuata(iscrizione, "mattino")
        ):
            db.session.rollback()
            return {
                "ok": False,
                "errore": "Prima devi scegliere il laboratorio del mattino.",
            }, 409

        campo_scelta = f"scelta_{tipologia}"
        campo_non_partecipa = f"non_partecipa_{tipologia}"
        scelta_precedente = getattr(iscrizione, campo_scelta) if iscrizione else None
        aveva_scelta = scelta_fascia_effettuata(iscrizione, tipologia)
        laboratori_da_bloccare = set()
        if laboratorio_id is not None:
            laboratori_da_bloccare.add(laboratorio_id)
        if scelta_precedente is not None:
            laboratori_da_bloccare.add(scelta_precedente)

        laboratori_bloccati = []
        if laboratori_da_bloccare:
            laboratori_bloccati = db.session.execute(
                db.select(Laboratorio)
                .where(Laboratorio.id.in_(laboratori_da_bloccare))
                .order_by(Laboratorio.id)
                .with_for_update()
            ).scalars().all()
        laboratorio = None
        if not non_partecipa:
            laboratorio = next(
                (item for item in laboratori_bloccati if item.id == laboratorio_id),
                None,
            )
            if laboratorio is None:
                db.session.rollback()
                return {"ok": False, "errore": "Laboratorio non valido."}, 400
            if laboratorio.tipologia != tipologia:
                db.session.rollback()
                return {
                    "ok": False,
                    "errore": "Il laboratorio appartiene a una fascia diversa.",
                }, 400

        if laboratorio is not None and scelta_precedente != laboratorio.id:
            colonna_scelta = getattr(Iscrizione, campo_scelta)
            occupati = Iscrizione.query.filter(
                colonna_scelta == laboratorio.id
            ).count()
            if occupati >= laboratorio.posti:
                db.session.rollback()
                return {
                    "ok": False,
                    "errore": "Il laboratorio si è appena riempito. Scegline un altro.",
                }, 409

        if iscrizione is None:
            iscrizione = Iscrizione(
                data=ora_roma(),
                partecipante=partecipante.id,
                scelta_mattino=None,
                scelta_pomeriggio=None,
                non_partecipa_mattino=False,
                non_partecipa_pomeriggio=False,
            )
            db.session.add(iscrizione)

        setattr(iscrizione, campo_scelta, laboratorio.id if laboratorio else None)
        setattr(iscrizione, campo_non_partecipa, non_partecipa)
        iscrizione.data = ora_roma()
        db.session.commit()

        completa = iscrizione_completa(iscrizione)
        if completa:
            destinazione = url_for("riepilogo_iscrizione")
        else:
            destinazione = url_for("scelta_laboratorio", tipologia="pomeriggio")
        if not aveva_scelta:
            messaggio = f"Scelta del sabato {tipologia} salvata."
        else:
            messaggio = f"Scelta del sabato {tipologia} modificata."
        flash(messaggio, "success")
        return {"ok": True, "messaggio": messaggio, "redirect": destinazione}
    except IntegrityError:
        db.session.rollback()
        return {
            "ok": False,
            "errore": "La richiesta è già stata elaborata. Ricarica la pagina.",
        }, 409
    except SQLAlchemyError:
        db.session.rollback()
        app.logger.exception("Errore durante il salvataggio del laboratorio")
        return {
            "ok": False,
            "errore": "Non è stato possibile salvare la scelta. Riprova.",
        }, 500

@app.route("/iscrizione/riepilogo")
def riepilogo_iscrizione():
    partecipante = get_partecipante_corrente()
    if partecipante is None:
        return redirect(url_for("index"))

    iscrizione = get_iscrizione_partecipante(partecipante.id)
    if (partecipante.deve_iscriversi_sabato
            and get_stato_iscrizioni() == "aperte" and not iscrizione_completa(iscrizione)):
        return redirect(url_for("percorso_iscrizione"))

    laboratorio_mattino = (
        db.session.get(Laboratorio, iscrizione.scelta_mattino)
        if iscrizione and iscrizione.scelta_mattino is not None
        else None
    )
    laboratorio_pomeriggio = (
        db.session.get(Laboratorio, iscrizione.scelta_pomeriggio)
        if iscrizione and iscrizione.scelta_pomeriggio is not None
        else None
    )
    return render_template(
        "riepilogo.html",
        partecipante=partecipante,
        iscrizione=iscrizione,
        laboratorio_mattino=laboratorio_mattino,
        laboratorio_pomeriggio=laboratorio_pomeriggio,
        non_partecipa_mattino=(
            iscrizione.non_partecipa_mattino if iscrizione else False
        ),
        non_partecipa_pomeriggio=(
            iscrizione.non_partecipa_pomeriggio if iscrizione else False
        ),
        completa=iscrizione_completa(iscrizione),
        nomi_gruppi_domenica=NOMI_GRUPPI_DOMENICA,
        iscrizioni_aperte=get_stato_iscrizioni() == "aperte",
    )

def riepilogo_confronto_import(operazioni):
    """Confronta record correnti e valori normalizzati senza modificare il DB."""
    risultato = dict(inseriti=0, aggiornati=0, invariati=0, totale=len(operazioni), dettaglio=[])
    for record, valori, etichetta, etichette_campi in operazioni:
        if record is None:
            risultato["inseriti"] += 1
            risultato["dettaglio"].append(etichetta + ": nuovo")
        else:
            cambiati = [campo for campo, valore in valori.items() if getattr(record, campo) != valore]
            risultato["aggiornati" if cambiati else "invariati"] += 1
            if cambiati:
                risultato["dettaglio"].append(etichetta + ": modificati " +
                    ", ".join(etichette_campi[c] for c in cambiati))
    return risultato


def confronta_partecipanti(dati, blocca=False):
    operazioni = []
    for riga in dati:
        query = Partecipante.query.filter_by(id=riga["id"]).populate_existing()
        record = (query.with_for_update() if blocca else query).first()
        operazioni.append((record, {c: v for c, v in riga.items() if c != "id"},
                           f"{riga['id']} – {riga['cognome']} {riga['nome']}", ETICHETTE_CAMPI_PARTECIPANTE))
    return operazioni, riepilogo_confronto_import(operazioni)


@app.route("/import_iscritti/valida", methods=["POST"])
@admin_required(api=True)
def valida_anteprima_partecipanti():
    dati, errore = valida_import_partecipanti(request.get_json(silent=True))
    if errore:
        return {"ok": False, "errore": errore}, 400
    try:
        _, riepilogo = confronta_partecipanti(dati)
        return {"ok": True, "partecipanti": dati, **riepilogo}
    except SQLAlchemyError:
        db.session.rollback()
        return {"ok": False, "errore": "Confronto non riuscito. Nessuna modifica salvata."}, 500


@app.route("/import_iscritti", methods=["GET", "POST"])
@admin_required()
def import_iscritti():
    if request.method == "POST":
        dati, errore = valida_import_partecipanti(request.get_json(silent=True))
        if errore:
            return {"ok": False, "errore": errore}, 400

        try:
            operazioni, riepilogo = confronta_partecipanti(dati, blocca=True)
            for riga, (partecipante, _, _, _) in zip(dati, operazioni):
                if partecipante is None:
                    partecipante = Partecipante(id=riga["id"])
                    db.session.add(partecipante)
                # Solo i campi autorizzati dal CSV: nessuna assegnazione o iscrizione.
                for campo, valore in riga.items():
                    if campo != "id":
                        setattr(partecipante, campo, valore)

            registra_ultimo_import(ULTIMO_IMPORT_PARTECIPANTI_KEY)
            ultimo_import = get_ultimo_import(ULTIMO_IMPORT_PARTECIPANTI_KEY)
            db.session.commit()
            return {
                "ok": True, **riepilogo,
                "ultimo_import": ultimo_import,
            }
        except SQLAlchemyError:
            db.session.rollback()
            app.logger.error("Errore durante l'import dei partecipanti")
            return {
                "ok": False,
                "errore": "Non è stato possibile completare l'import dei partecipanti. Nessuna modifica salvata.",
            }, 500
    return render_template(
        "import_iscritti.html",
        ultimo_import=get_ultimo_import(ULTIMO_IMPORT_PARTECIPANTI_KEY),
        colonne_csv=COLONNE_CSV_PARTECIPANTI,
    )

@app.route("/import_laboratori/valida", methods=["POST"], defaults={"anteprima": True})
@app.route("/import_laboratori", methods=["GET", "POST"])
@admin_required()
def import_laboratori(anteprima=False):
    if request.method == "POST":
        dati, errore = valida_import_laboratori(request.get_json(silent=True))
        if errore:
            return {"ok": False, "errore": errore}, 400

        try:
            # Stesso ordine dei lock usato dal salvataggio delle iscrizioni.
            # Su MariaDB restano acquisiti fino al commit/rollback: i conteggi
            # successivi non possono essere superati da nuove assegnazioni.
            if not anteprima:
                Laboratorio.query.order_by(Laboratorio.id).with_for_update().populate_existing().all()
            operazioni = []
            id_visti = set()
            for tipologia, chiave in (
                ("mattino", "lab_mattino"),
                ("pomeriggio", "lab_pomeriggio"),
            ):
                for riga in dati[chiave]:
                    query = Laboratorio.query.filter_by(
                        tipologia=tipologia, id_lab=riga["id"],
                    ).order_by(Laboratorio.id).populate_existing()
                    corrispondenze = (query if anteprima else query.with_for_update()).all()
                    if len(corrispondenze) > 1:
                        db.session.rollback()
                        return {
                            "ok": False,
                            "errore": f"Import annullato: più laboratori nel database per ({tipologia}, {riga['id']}). Risolvi i duplicati prima di reimportare.",
                        }, 409
                    laboratorio = corrispondenze[0] if corrispondenze else None
                    if laboratorio is not None:
                        # Anche codici distinti nel JSON possono coincidere
                        # secondo la collation del database.
                        if laboratorio.id in id_visti:
                            db.session.rollback()
                            return {
                                "ok": False,
                                "errore": f"Import annullato: più righe identificano il laboratorio {riga['id']} del {tipologia}.",
                            }, 400
                        id_visti.add(laboratorio.id)
                        if riga["posti"] < laboratorio.posti:
                            scelta = getattr(Iscrizione, f"scelta_{tipologia}")
                            occupati = Iscrizione.query.filter(scelta == laboratorio.id).count()
                            if riga["posti"] < occupati:
                                db.session.rollback()
                                return {
                                    "ok": False,
                                    "errore": f"Import annullato: il laboratorio {riga['id']} del {tipologia} ha {occupati} iscritti; non puoi ridurre i posti a {riga['posti']}.",
                                }, 409
                    operazioni.append((tipologia, riga, laboratorio))

            riepilogo = riepilogo_confronto_import([
                (laboratorio, {c: riga[c] for c in ("titolo", "descrizione", "posti")},
                 f"{tipologia} / {riga['id']} – {riga['titolo']}",
                 {"titolo": "Titolo", "descrizione": "Descrizione", "posti": "Posti"})
                for tipologia, riga, laboratorio in operazioni
            ])
            if anteprima:
                return {"ok": True, **riepilogo}
            # Nessuna scrittura prima della verifica di entrambe le fasce.
            for tipologia, riga, laboratorio in operazioni:
                if laboratorio is None:
                    laboratorio = Laboratorio(id_lab=riga["id"], tipologia=tipologia)
                    db.session.add(laboratorio)
                laboratorio.titolo = riga["titolo"]
                laboratorio.descrizione = riga["descrizione"]
                laboratorio.posti = riga["posti"]

            registra_ultimo_import(ULTIMO_IMPORT_LABORATORI_KEY)
            ultimo_import = get_ultimo_import(ULTIMO_IMPORT_LABORATORI_KEY)
            db.session.commit()
            return {
                "ok": True, **riepilogo, "ultimo_import": ultimo_import,
            }
        except SQLAlchemyError:
            db.session.rollback()
            app.logger.exception("Errore durante l'import dei laboratori")
            return {
                "ok": False,
                "errore": "Non è stato possibile completare l'import dei laboratori. Nessuna modifica salvata.",
            }, 500
    return render_template(
        "import_lab.html",
        ultimo_import=get_ultimo_import(ULTIMO_IMPORT_LABORATORI_KEY),
    )

def intero_form(valore):
    testo = str(valore or "").strip()
    return int(testo) if re.fullmatch(r"[0-9]{1,10}", testo) else None


@app.route("/admin/partecipanti/nuovo", methods=["GET", "POST"])
@app.route("/admin/partecipanti/<int:partecipante_id>/modifica", methods=["GET", "POST"])
@admin_required()
def modifica_partecipante(partecipante_id=None):
    persona = db.session.get(Partecipante, partecipante_id) if partecipante_id is not None else None
    if partecipante_id is not None and persona is None:
        abort(404)
    valori = {campo: getattr(persona, campo) if persona else ""
              for campo in COLONNE_CSV_PARTECIPANTI.values()}
    iscrizione = get_iscrizione_partecipante(partecipante_id) if persona is not None else None
    assegnazioni = {"gruppo_domenica": str(persona.gruppo_domenica or "") if persona else ""}
    for fascia in ("mattino", "pomeriggio"):
        assegnazioni[f"scelta_{fascia}"] = ("non_partecipa" if iscrizione and getattr(iscrizione, f"non_partecipa_{fascia}")
            else str(getattr(iscrizione, f"scelta_{fascia}", None) or ""))
        assegnazioni[f"sottogruppo_{fascia}"] = getattr(iscrizione, f"sottogruppo_{fascia}", None) or ""
    if request.method == "POST":
        if persona is not None:
            assegnazioni = {campo: request.form.get(campo, valore) for campo, valore in assegnazioni.items()}
        valori = {campo: request.form.get(campo, "") for campo in COLONNE_CSV_PARTECIPANTI.values()}
        valori["id"] = intero_form(valori["id"])
        dati, errore = valida_import_partecipanti([valori])
        if persona is not None and valori["id"] != persona.id:
            errore = "Il codice censimento non può essere modificato."
        if not errore:
            for campo in CAMPI_MANUALI_OBBLIGATORI:
                if not dati[0][campo]:
                    errore = f"Il campo {ETICHETTE_CAMPI_PARTECIPANTE[campo]} è obbligatorio."
                    break
            if not errore and not re.fullmatch(r"[^\s@]+@[^\s@]+\.[^\s@]+", dati[0]["email"]):
                errore = "Inserisci un indirizzo Email valido."
            if not errore and "altro" in (dati[0]["ruolo"] or "").casefold() and not dati[0]["incarico_altro"]:
                errore = "Specifica Incarico altro quando il ruolo indica Altro."
        if not errore and persona is not None:
            gruppo = assegnazioni["gruppo_domenica"]
            if gruppo and (intero_form(gruppo) not in NOMI_GRUPPI_DOMENICA):
                errore = "Gruppo domenica non valido."
            for fascia in ("mattino", "pomeriggio"):
                scelta = assegnazioni[f"scelta_{fascia}"]
                if scelta not in ("", "non_partecipa"):
                    lab = db.session.get(Laboratorio, intero_form(scelta)) if intero_form(scelta) else None
                    if lab is None or lab.tipologia != fascia:
                        errore = f"Laboratorio {fascia} non valido."
                if assegnazioni[f"sottogruppo_{fascia}"] not in ("", "A", "B"):
                    errore = f"Sottogruppo {fascia} non valido."
        if not errore:
            try:
                if persona is None:
                    if db.session.get(Partecipante, valori["id"]) is not None:
                        errore = "Questo codice censimento è già presente."
                    else:
                        persona = Partecipante(id=valori["id"], gruppo_domenica=None)
                        db.session.add(persona)
                else:
                    persona = Partecipante.query.filter_by(id=partecipante_id).with_for_update().one()
                if not errore:
                    for campo, valore in dati[0].items():
                        if campo != "id":
                            setattr(persona, campo, valore)
                    if partecipante_id is not None and any(c in request.form for c in assegnazioni):
                        # Correzione forzata: lock coerenti con il flusso utente, senza limite capienza.
                        ids = [intero_form(assegnazioni[f"scelta_{f}"]) for f in ("mattino", "pomeriggio")]
                        if iscrizione:
                            ids += [iscrizione.scelta_mattino, iscrizione.scelta_pomeriggio]
                        Laboratorio.query.filter(Laboratorio.id.in_([i for i in ids if i])).order_by(Laboratorio.id).with_for_update().all()
                        iscrizione = get_iscrizione_partecipante(partecipante_id)
                        if iscrizione is None and any(assegnazioni[f"scelta_{f}"] for f in ("mattino", "pomeriggio")):
                            iscrizione = Iscrizione(partecipante=partecipante_id, data=ora_roma())
                            db.session.add(iscrizione)
                        if iscrizione is not None:
                            for fascia in ("mattino", "pomeriggio"):
                                scelta = assegnazioni[f"scelta_{fascia}"]
                                laboratorio_id = intero_form(scelta)
                                setattr(iscrizione, f"scelta_{fascia}", laboratorio_id)
                                setattr(iscrizione, f"non_partecipa_{fascia}", scelta == "non_partecipa")
                                setattr(iscrizione, f"sottogruppo_{fascia}", (assegnazioni[f"sottogruppo_{fascia}"] or None) if laboratorio_id else None)
                                if laboratorio_id:
                                    persona.deve_iscriversi_sabato = True
                        persona.gruppo_domenica = intero_form(assegnazioni["gruppo_domenica"])
                        if persona.gruppo_domenica:
                            persona.includi_domenica = True
                    db.session.commit()
                    flash("Partecipante salvato.", "success")
                    return redirect(url_for("gestione_iscrizioni"))
            except IntegrityError:
                db.session.rollback()
                errore = "Salvataggio non riuscito: verifica che il codice censimento non sia già presente."
            except SQLAlchemyError:
                db.session.rollback()
                errore = "Non è stato possibile salvare il partecipante. Nessuna modifica salvata."
        db.session.rollback()
        flash(errore, "warning")
        if partecipante_id is not None:
            valori["id"] = partecipante_id
    return render_template("form_partecipante.html", valori=valori,
                           modifica=partecipante_id is not None,
                           colonne=COLONNE_CSV_PARTECIPANTI, obbligatori=CAMPI_MANUALI_OBBLIGATORI,
                           assegnazioni=assegnazioni, nomi_gruppi=NOMI_GRUPPI_DOMENICA,
                           laboratori=Laboratorio.query.order_by(Laboratorio.tipologia, Laboratorio.id_lab).all()), (400 if request.method == "POST" else 200)


@app.route("/admin/laboratori/nuovo", methods=["GET", "POST"])
@app.route("/admin/laboratori/<int:laboratorio_id>/modifica", methods=["GET", "POST"])
@admin_required()
def modifica_laboratorio(laboratorio_id=None):
    laboratorio = db.session.get(Laboratorio, laboratorio_id) if laboratorio_id is not None else None
    if laboratorio_id is not None and laboratorio is None:
        abort(404)
    valori = {campo: getattr(laboratorio, campo) if laboratorio else ""
              for campo in ("tipologia", "id_lab", "titolo", "descrizione", "posti")}
    if request.method == "POST":
        valori = {campo: request.form.get(campo, "") for campo in valori}
        riga, errore = valida_laboratorio({
            "id": valori["id_lab"], "titolo": valori["titolo"],
            "descrizione": valori["descrizione"], "posti": intero_form(valori["posti"]),
        })
        if valori["tipologia"] not in ("mattino", "pomeriggio"):
            errore = "Scegli mattino o pomeriggio."
        if laboratorio is not None and (valori["tipologia"] != laboratorio.tipologia
                                        or valori["id_lab"] != laboratorio.id_lab):
            errore = "Codice e fascia del laboratorio non possono essere modificati."
        if not errore:
            try:
                if laboratorio is None:
                    if Laboratorio.query.filter_by(tipologia=valori["tipologia"], id_lab=riga["id"]).first():
                        errore = "Esiste già un laboratorio con questo codice nella fascia selezionata."
                    else:
                        laboratorio = Laboratorio(id_lab=riga["id"], tipologia=valori["tipologia"])
                        db.session.add(laboratorio)
                else:
                    laboratorio = Laboratorio.query.filter_by(id=laboratorio_id).populate_existing().with_for_update().one()
                    scelta = getattr(Iscrizione, f"scelta_{laboratorio.tipologia}")
                    occupati = Iscrizione.query.filter(scelta == laboratorio.id).count()
                    if riga["posti"] < occupati:
                        errore = f"Il laboratorio ha {occupati} iscritti: i posti non possono essere inferiori."
                if not errore:
                    for campo in ("titolo", "descrizione", "posti"):
                        setattr(laboratorio, campo, riga[campo])
                    db.session.commit()
                    flash("Laboratorio salvato.", "success")
                    return redirect(url_for("gestione_dati"))
            except SQLAlchemyError:
                db.session.rollback()
                errore = "Non è stato possibile salvare il laboratorio. Nessuna modifica salvata."
        db.session.rollback()
        flash(errore, "warning")
        if laboratorio_id is not None:
            valori["id_lab"] = laboratorio.id_lab
            valori["tipologia"] = laboratorio.tipologia
    return render_template("form_laboratorio.html", valori=valori,
                           modifica=laboratorio_id is not None), (400 if request.method == "POST" else 200)


@app.route("/admin/partecipanti/<int:partecipante_id>/reset_sabato", methods=["POST"])
@admin_required()
def reset_sabato_partecipante(partecipante_id):
    try:
        persona = Partecipante.query.filter_by(id=partecipante_id).with_for_update().first()
        if persona is None:
            abort(404)
        iscrizione = get_iscrizione_partecipante(partecipante_id)
        if iscrizione is not None:
            laboratori_id = [id_lab for id_lab in (iscrizione.scelta_mattino, iscrizione.scelta_pomeriggio)
                             if id_lab is not None]
            Laboratorio.query.filter(Laboratorio.id.in_(laboratori_id)).order_by(Laboratorio.id).with_for_update().all()
            db.session.delete(iscrizione)
        db.session.commit()
        flash(f"Iscrizioni del sabato azzerate per il codice {partecipante_id}.", "success")
    except SQLAlchemyError:
        db.session.rollback()
        flash("Non è stato possibile azzerare le iscrizioni. Nessuna modifica salvata.", "warning")
    return redirect(url_for("gestione_iscrizioni"))


def errore_richiesta_reset(conferma_attesa):
    password_attuale = request.form.get("password_attuale", "")
    if not check_password_hash(current_user.password, password_attuale):
        return "Password non corretta. Nessun dato è stato cancellato."
    if request.form.get("conferma") != conferma_attesa:
        return (
            "Conferma esplicitamente l'operazione. "
            "Nessun dato è stato cancellato."
        )
    return None

@app.route("/admin/gestione_dati")
@admin_required()
def gestione_dati():
    return render_template("gestione_dati.html", laboratori=Laboratorio.query.order_by(Laboratorio.tipologia, Laboratorio.id_lab, Laboratorio.id).all())

@app.route("/admin/gestione_dati/reset_iscrizioni", methods=["POST"])
@admin_required()
def reset_iscrizioni():
    errore = errore_richiesta_reset("iscrizioni")
    if errore:
        flash(errore, "warning")
        return redirect(url_for("gestione_dati"))

    try:
        Iscrizione.query.delete()
        db.session.commit()
    except SQLAlchemyError:
        db.session.rollback()
        app.logger.exception("Errore durante il reset delle iscrizioni")
        flash(
            "Non è stato possibile cancellare le iscrizioni. "
            "Nessun dato è stato cancellato.",
            "danger",
        )
        return redirect(url_for("gestione_dati"))

    flash(
        "Tutte le iscrizioni registrate sono state cancellate.",
        "success",
    )
    return redirect(url_for("gestione_dati"))

@app.route("/admin/gestione_dati/reset_partecipanti", methods=["POST"])
@admin_required()
def reset_partecipanti():
    errore = errore_richiesta_reset("partecipanti")
    if errore:
        flash(errore, "warning")
        return redirect(url_for("gestione_dati"))

    try:
        if Iscrizione.query.first() is not None:
            flash(
                "Non è possibile cancellare i partecipanti perché esistono "
                "iscrizioni registrate. Esegui prima il reset delle iscrizioni.",
                "warning",
            )
            return redirect(url_for("gestione_dati"))

        Partecipante.query.delete()
        SysOption.query.filter_by(key=ULTIMO_IMPORT_PARTECIPANTI_KEY).delete()
        db.session.commit()
    except SQLAlchemyError:
        db.session.rollback()
        app.logger.exception("Errore durante il reset dei partecipanti")
        flash(
            "Non è stato possibile cancellare i partecipanti. "
            "Nessun dato è stato cancellato.",
            "danger",
        )
        return redirect(url_for("gestione_dati"))

    flash(
        "Tutti i partecipanti importati sono stati cancellati.",
        "success",
    )
    return redirect(url_for("gestione_dati"))

@app.route("/admin/gestione_dati/reset_laboratori", methods=["POST"])
@admin_required()
def reset_laboratori():
    errore = errore_richiesta_reset("laboratori")
    if errore:
        flash(errore, "warning")
        return redirect(url_for("gestione_dati"))

    try:
        if Iscrizione.query.first() is not None:
            flash(
                "Non è possibile cancellare i laboratori perché esistono "
                "iscrizioni registrate. Esegui prima il reset delle iscrizioni.",
                "warning",
            )
            return redirect(url_for("gestione_dati"))

        Laboratorio.query.delete()
        SysOption.query.filter_by(key=ULTIMO_IMPORT_LABORATORI_KEY).delete()
        db.session.commit()
    except SQLAlchemyError:
        db.session.rollback()
        app.logger.exception("Errore durante il reset dei laboratori")
        flash(
            "Non è stato possibile cancellare i laboratori. "
            "Nessun dato è stato cancellato.",
            "danger",
        )
        return redirect(url_for("gestione_dati"))

    flash(
        "Tutti i laboratori importati sono stati cancellati.",
        "success",
    )
    return redirect(url_for("gestione_dati"))

@app.route("/admin/stato_iscrizioni", methods=["GET", "POST"])
@admin_required()
def stato_iscrizioni():
    if request.method == "POST":
        nuovo_stato = request.form.get("stato")
        if nuovo_stato not in STATI_ISCRIZIONI_VALIDI:
            flash("Stato iscrizioni non valido.", "warning")
            return render_template(
                "stato_iscrizioni.html",
                stato_corrente=get_stato_iscrizioni(),
                stati_validi=STATI_ISCRIZIONI_VALIDI,
                messaggio_iscrizioni=get_messaggio_iscrizioni(),
            ), 400

        messaggio_iscrizioni = request.form.get("messaggio_iscrizioni", "").strip()
        if len(messaggio_iscrizioni) > 128:
            flash("Il messaggio non può superare 128 caratteri.", "warning")
            return render_template(
                "stato_iscrizioni.html",
                stato_corrente=get_stato_iscrizioni(),
                stati_validi=STATI_ISCRIZIONI_VALIDI,
                messaggio_iscrizioni=get_messaggio_iscrizioni(),
            ), 400

        set_sys_option(STATO_ISCRIZIONI_KEY, nuovo_stato)
        set_sys_option(MESSAGGIO_ISCRIZIONI_KEY, messaggio_iscrizioni)

        db.session.commit()
        flash("Stato iscrizioni salvato correttamente.", "success")
        return redirect(url_for("stato_iscrizioni"))

    return render_template(
        "stato_iscrizioni.html",
        stato_corrente=get_stato_iscrizioni(),
        stati_validi=STATI_ISCRIZIONI_VALIDI,
        messaggio_iscrizioni=get_messaggio_iscrizioni(),
    )

@app.route("/admin/iscrizioni")
@admin_required()
def gestione_iscrizioni():
    laboratorio_mattino = aliased(Laboratorio)
    laboratorio_pomeriggio = aliased(Laboratorio)
    iscrizioni = (
        db.session.query(
            Iscrizione,
            Partecipante,
            laboratorio_mattino,
            laboratorio_pomeriggio,
        )
        .select_from(Partecipante)
        .outerjoin(Iscrizione, Iscrizione.partecipante == Partecipante.id)
        .outerjoin(
            laboratorio_mattino,
            Iscrizione.scelta_mattino == laboratorio_mattino.id,
        )
        .outerjoin(
            laboratorio_pomeriggio,
            Iscrizione.scelta_pomeriggio == laboratorio_pomeriggio.id,
        )
        .order_by(Partecipante.cognome, Partecipante.nome)
        .all()
    )

    iscrizioni = [
        (*riga, stato_iscrizione(riga[0], riga[1]))
        for riga in iscrizioni
    ]
    totale_partecipanti = len(iscrizioni)
    totale_sabato = sum(p.deve_iscriversi_sabato for _, p, *_ in iscrizioni)
    iscrizioni_complete = sum(
        stato == "completo" for *_, stato in iscrizioni
    )
    iscrizioni_incomplete = sum(
        stato == "incompleto" for *_, stato in iscrizioni
    )
    partecipanti_non_iniziati = sum(
        stato == "non_iniziato" for *_, stato in iscrizioni
    )
    percentuale_completamento = (
        (iscrizioni_complete / totale_sabato * 100)
        if totale_sabato
        else 0
    )

    return render_template(
        "gestione_iscrizioni.html",
        iscrizioni=iscrizioni,
        totale_partecipanti=totale_partecipanti,
        totale_sabato=totale_sabato,
        iscrizioni_complete=iscrizioni_complete,
        iscrizioni_incomplete=iscrizioni_incomplete,
        partecipanti_non_iniziati=partecipanti_non_iniziati,
        percentuale_completamento=percentuale_completamento,
    )

def leggi_distribuzione_ab():
    """Riepilogo delle iscrizioni effettive e delle A/B già persistite."""
    conteggi = {}
    for fascia in ("mattino", "pomeriggio"):
        scelta = getattr(Iscrizione, f"scelta_{fascia}")
        sottogruppo = getattr(Iscrizione, f"sottogruppo_{fascia}")
        righe = db.session.query(scelta, sottogruppo, func.count(Iscrizione.id)).filter(
            scelta.isnot(None), getattr(Iscrizione, f"non_partecipa_{fascia}").is_(False),
        ).group_by(scelta, sottogruppo).all()
        for laboratorio_id, gruppo, totale in righe:
            valori = conteggi.setdefault((fascia, laboratorio_id), {"totale": 0, "A": 0, "B": 0})
            valori["totale"] += totale
            if gruppo in ("A", "B"):
                valori[gruppo] += totale
    return [
        {"laboratorio": laboratorio,
         **conteggi.get((laboratorio.tipologia, laboratorio.id), {"totale": 0, "A": 0, "B": 0})}
        for laboratorio in Laboratorio.query.filter(
            Laboratorio.tipologia.in_(("mattino", "pomeriggio")),
        ).order_by(Laboratorio.tipologia, Laboratorio.id).all()
    ]


def ricalcola_sottogruppi_ab():
    try:
        # Stesso ordine del salvataggio scelte: partecipanti, laboratori, iscrizioni.
        partecipanti = (Partecipante.query.order_by(Partecipante.id)
                        .populate_existing().with_for_update().all())
        laboratori = (Laboratorio.query.filter(Laboratorio.tipologia.in_(("mattino", "pomeriggio")))
                      .order_by(Laboratorio.id).populate_existing().with_for_update().all())
        iscrizioni = (Iscrizione.query.order_by(Iscrizione.id)
                      .populate_existing().with_for_update().all())
        persone = {persona.id: persona for persona in partecipanti}
        for iscrizione in iscrizioni:
            iscrizione.sottogruppo_mattino = None
            iscrizione.sottogruppo_pomeriggio = None
        for fascia in ("mattino", "pomeriggio"):
            for laboratorio in laboratori:
                if laboratorio.tipologia != fascia:
                    continue
                iscritti = [i for i in iscrizioni
                            if getattr(i, f"scelta_{fascia}") == laboratorio.id
                            and not getattr(i, f"non_partecipa_{fascia}")]
                assegnazioni = calcola_sottogruppi_ab(persone[i.partecipante] for i in iscritti)
                for iscrizione in iscritti:
                    setattr(iscrizione, f"sottogruppo_{fascia}", assegnazioni[iscrizione.partecipante])
        db.session.commit()
    except Exception:
        db.session.rollback()
        raise


@app.route("/admin/iscrizioni/gruppi-ab")
@admin_required()
def gruppi_ab_sabato():
    return render_template("gruppi_ab_sabato.html", distribuzione=leggi_distribuzione_ab(), iscrizioni_aperte=get_stato_iscrizioni() == "aperte")


@app.route("/admin/iscrizioni/gruppi-ab/ricalcola", methods=["POST"])
@admin_required(api=True)
def ricalcola_ab_sabato():
    if get_stato_iscrizioni() == "aperte":
        return {"ok": False, "errore": "Chiudi le iscrizioni prima di calcolare i gruppi A/B."}, 409
    if request.form.get("conferma") != "ricalcola":
        return {"ok": False, "errore": "Conferma il ricalcolo dei gruppi A/B sabato."}, 400
    try:
        ricalcola_sottogruppi_ab()
    except Exception:
        app.logger.exception("Errore durante il ricalcolo dei gruppi A/B sabato")
        flash("Ricalcolo non riuscito. Tutte le assegnazioni A/B precedenti sono state conservate.", "danger")
        return render_template("gruppi_ab_sabato.html", distribuzione=leggi_distribuzione_ab(), iscrizioni_aperte=get_stato_iscrizioni() == "aperte"), 500
    flash("Gruppi A/B sabato ricalcolati per tutti i laboratori.", "success")
    return redirect(url_for("gruppi_ab_sabato"))


def formatta_foglio_excel(foglio, larghezze_colonne):
    foglio.freeze_panes = "A2"
    foglio.auto_filter.ref = foglio.dimensions
    for colonna, larghezza in zip(foglio.columns, larghezze_colonne):
        foglio.column_dimensions[colonna[0].column_letter].width = larghezza


def invia_file_excel(workbook, prefisso_nome_file):
    file_excel = io.BytesIO()
    workbook.save(file_excel)
    file_excel.seek(0)
    data_esportazione = datetime.now(ZoneInfo("Europe/Rome")).date().isoformat()
    nome_file = f"{prefisso_nome_file}_{data_esportazione}.xlsx"

    return send_file(
        file_excel,
        as_attachment=True,
        download_name=nome_file,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


def nome_foglio_univoco(nome_base, nomi_usati):
    nome_base = " ".join(nome_base.split())
    nome_base = re.sub(r"[\\/*?:\[\]]", "-", nome_base)
    nome = nome_base[:31]
    numero = 2
    while nome.casefold() in nomi_usati:
        suffisso = f" ({numero})"
        nome = f"{nome_base[:31 - len(suffisso)]}{suffisso}"
        numero += 1
    nomi_usati.add(nome.casefold())
    return nome


def nome_foglio_laboratorio(laboratorio, nomi_usati):
    sigla_tipologia = {
        "mattino": "M",
        "pomeriggio": "P",
    }.get(laboratorio.tipologia, "L")
    return nome_foglio_univoco(
        f"{sigla_tipologia} - {laboratorio.id_lab} - {laboratorio.titolo}",
        nomi_usati,
    )


def dati_export_iscrizioni():
    return (db.session.query(Partecipante, Iscrizione)
            .outerjoin(Iscrizione, Iscrizione.partecipante == Partecipante.id)
            .order_by(Partecipante.cognome, Partecipante.nome, Partecipante.id).all())


def celle_export_fascia(partecipante, iscrizione, fascia, laboratori):
    if not partecipante.deve_iscriversi_sabato:
        return "Iscrizione non richiesta", "Iscrizione non richiesta"
    if iscrizione and getattr(iscrizione, f"non_partecipa_{fascia}"):
        return TESTO_NON_PARTECIPA, TESTO_NON_PARTECIPA
    laboratorio = laboratori.get(getattr(iscrizione, f"scelta_{fascia}", None))
    if laboratorio is None:
        return "Non iscritto", "Non iscritto"
    return (f"{laboratorio.id_lab} - {laboratorio.titolo}",
            getattr(iscrizione, f"sottogruppo_{fascia}") or "Non assegnato")


def compila_foglio_anagrafica(foglio, partecipanti, colonne_extra=(), valori_extra=None):
    foglio.append(["Codice censimento", "Nome", "Cognome", "Gruppo", "Zona", "Regione", "FoCa", "Email"] + list(colonne_extra))
    for persona in partecipanti:
        foglio.append([persona.id, persona.nome, persona.cognome, persona.gruppo,
                       persona.zona, persona.regione, persona.foca, persona.email]
                      + (list(valori_extra(persona)) if valori_extra else []))
    formatta_foglio_excel(foglio, (20, 24, 24, 24, 24, 24, 24, 36) + (32,) * len(colonne_extra))


@app.route("/admin/iscrizioni/esporta/laboratori")
@admin_required()
def scarica_iscrizioni_per_laboratorio():
    laboratori = Laboratorio.query.order_by(
        Laboratorio.tipologia, Laboratorio.id_lab, Laboratorio.titolo, Laboratorio.id,
    ).all()
    partecipanti = dati_export_iscrizioni()
    iscritti_per_laboratorio = {}
    non_partecipanti = {"mattino": [], "pomeriggio": []}
    for partecipante, iscrizione in partecipanti:
        if iscrizione is None or not partecipante.deve_iscriversi_sabato:
            continue
        for fascia in ("mattino", "pomeriggio"):
            if getattr(iscrizione, f"non_partecipa_{fascia}"):
                non_partecipanti[fascia].append(partecipante)
            else:
                laboratorio_id = getattr(iscrizione, f"scelta_{fascia}")
                if laboratorio_id is not None:
                    iscritti_per_laboratorio.setdefault((fascia, laboratorio_id), []).append(
                        partecipante
                    )

    workbook = Workbook()
    foglio = workbook.active
    foglio.title = "Tutti i partecipanti"
    iscrizioni = {p.id: i for p, i in partecipanti}
    laboratori_per_id = {lab.id: lab for lab in laboratori}

    def situazione_sabato(persona):
        valori = []
        for fascia in ("mattino", "pomeriggio"):
            stato, ab = celle_export_fascia(persona, iscrizioni[persona.id], fascia, laboratori_per_id)
            # Negli elenchi sabato esportare solo A/B assegnati; le altre celle restano vuote.
            valori.extend((stato, ab if ab in ("A", "B") else None))
        return valori

    compila_foglio_anagrafica(foglio, (p for p, _ in partecipanti),
                             ("Laboratorio mattino", "Gruppo A/B mattino",
                              "Laboratorio pomeriggio", "Gruppo A/B pomeriggio"), situazione_sabato)
    nomi_usati = {foglio.title.casefold()}
    for laboratorio in laboratori:
        foglio = workbook.create_sheet(nome_foglio_laboratorio(laboratorio, nomi_usati))
        compila_foglio_anagrafica(foglio, iscritti_per_laboratorio.get((laboratorio.tipologia, laboratorio.id), []),
                                 ("Gruppo A/B",),
                                 lambda p: (getattr(iscrizioni[p.id], f"sottogruppo_{laboratorio.tipologia}"),))

    for tipologia, nome_foglio in (("mattino", "M - Non partecipa"), ("pomeriggio", "P - Non partecipa")):
        foglio = workbook.create_sheet(nome_foglio_univoco(nome_foglio, nomi_usati))
        compila_foglio_anagrafica(foglio, non_partecipanti[tipologia])
    return invia_file_excel(workbook, "iscrizioni_per_laboratorio")


def errore_api_suddivisione(messaggio, stato=400):
    return {"ok": False, "errore": messaggio}, stato


def leggi_distribuzione_domenica():
    """Legge esclusivamente le assegnazioni persistite, compresi i gruppi vuoti."""
    conteggi = dict(db.session.query(
        Partecipante.gruppo_domenica, func.count(Partecipante.id),
    ).filter(Partecipante.gruppo_domenica.isnot(None)).group_by(
        Partecipante.gruppo_domenica,
    ).all())
    gruppi = {numero: conteggi.get(numero, 0)
              for numero in range(1, NUMERO_GRUPPI_DOMENICA + 1)}
    return {
        "gruppi": gruppi,
        "totale_assegnati": sum(gruppi.values()),
        "dimensione_minima": min(gruppi.values()),
        "dimensione_massima": max(gruppi.values()),
    }


def ricalcola_gruppi_domenica():
    try:
        partecipanti = (Partecipante.query.order_by(Partecipante.id)
                        .populate_existing().with_for_update().all())
        assegnazioni = calcola_gruppi_domenica(partecipanti)
        for persona in partecipanti:
            persona.gruppo_domenica = assegnazioni.get(persona.id)
        set_sys_option(ULTIMO_RICALCOLO_DOMENICA_KEY, datetime.now(ZoneInfo("Europe/Rome")).isoformat())
        db.session.commit()
        return len(assegnazioni)
    except Exception:
        db.session.rollback()
        raise


@app.route("/admin/suddivisione-gruppi")
@admin_required()
def suddivisione_gruppi():
    return pagina_suddivisione_domenica()


def pagina_suddivisione_domenica():
    return render_template(
        "suddivisione_gruppi.html",
        inclusi=Partecipante.query.filter(Partecipante.includi_domenica.is_(True)).count(),
        numero_gruppi=NUMERO_GRUPPI_DOMENICA,
        nomi_gruppi_domenica=NOMI_GRUPPI_DOMENICA,
        riepilogo=leggi_distribuzione_domenica(),
        ultimo_ricalcolo=get_ultimo_import(ULTIMO_RICALCOLO_DOMENICA_KEY),
    )


@app.route("/admin/suddivisione-gruppi/esporta")
@admin_required()
def esporta_gruppi_domenica():
    partecipanti = (Partecipante.query.filter(Partecipante.includi_domenica.is_(True))
                    .order_by(Partecipante.cognome, Partecipante.nome, Partecipante.id).all())
    workbook = Workbook()
    generale = workbook.active
    generale.title = "Tutti i partecipanti"
    compila_foglio_anagrafica(generale, partecipanti, ("Gruppo domenica",),
                             lambda p: (NOMI_GRUPPI_DOMENICA.get(p.gruppo_domenica, "Non assegnato"),))
    for numero, nome in NOMI_GRUPPI_DOMENICA.items():
        foglio = workbook.create_sheet(nome)
        compila_foglio_anagrafica(foglio, (p for p in partecipanti if p.gruppo_domenica == numero))
    return invia_file_excel(workbook, "gruppi_domenica")


@app.route("/admin/suddivisione-gruppi/ricalcola", methods=["POST"])
@admin_required(api=True)
def ricalcola_domenica():
    if request.form.get("conferma") != "ricalcola":
        return errore_api_suddivisione("Conferma il ricalcolo dei gruppi domenica.")
    try:
        totale = ricalcola_gruppi_domenica()
    except Exception:
        app.logger.exception("Errore durante il ricalcolo dei gruppi domenica")
        flash("Ricalcolo non riuscito. Le assegnazioni precedenti sono state conservate.", "danger")
        return pagina_suddivisione_domenica(), 500
    flash(f"Gruppi domenica ricalcolati: {totale} partecipanti assegnati.", "success")
    return redirect(url_for("suddivisione_gruppi"))


@app.route("/admin/suddivisione-gruppi/valida", methods=["POST"])
@admin_required(api=True)
def valida_excel_suddivisione():
    workbook = None
    try:
        workbook = carica_excel_suddivisione(request.files.get("file"))
        _, _, partecipanti = estrai_partecipanti_suddivisione(workbook)
        return {"ok": True, "partecipanti": len(partecipanti)}
    except ErroreSuddivisione as errore:
        return errore_api_suddivisione(str(errore))
    except HTTPException:
        raise
    except Exception:
        app.logger.exception("Errore durante la validazione del file di suddivisione")
        return errore_api_suddivisione(
            "Non è stato possibile verificare il file Excel.",
            500,
        )
    finally:
        if workbook is not None:
            workbook.close()


@app.route("/admin/suddivisione-gruppi/genera", methods=["POST"])
@admin_required(api=True)
def genera_suddivisione_gruppi():
    workbook = None
    try:
        workbook = carica_excel_suddivisione(request.files.get("file"))
        _, _, partecipanti = estrai_partecipanti_suddivisione(workbook)
        numero_gruppi = valida_numero_gruppi(
            request.form.get("numero_laboratori"),
            len(partecipanti),
        )
        gruppi = genera_excel_suddivisione(
            workbook,
            partecipanti,
            numero_gruppi,
        )
        file_excel = io.BytesIO()
        workbook.save(file_excel)
        file_excel.seek(0)
        data_esportazione = datetime.now(ZoneInfo("Europe/Rome")).date().isoformat()
        risposta = send_file(
            file_excel,
            as_attachment=True,
            download_name=f"suddivisione_laboratori_{data_esportazione}.xlsx",
            mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        conteggio_dimensioni = distribuzione_dimensioni(gruppi)
        risposta.headers["X-Partecipanti"] = str(len(partecipanti))
        risposta.headers["X-Laboratori"] = str(numero_gruppi)
        risposta.headers["X-Dimensioni-Gruppi"] = json.dumps(
            conteggio_dimensioni,
            separators=(",", ":"),
            sort_keys=True,
        )
        return risposta
    except ErroreSuddivisione as errore:
        return errore_api_suddivisione(str(errore))
    except HTTPException:
        raise
    except Exception:
        app.logger.exception("Errore durante la generazione della suddivisione")
        return errore_api_suddivisione(
            "Non è stato possibile generare il file Excel.",
            500,
        )
    finally:
        if workbook is not None:
            workbook.close()


def leggi_dati_comunicazioni():
    laboratori = {lab.id: lab for lab in Laboratorio.query.all()}
    righe = []
    for partecipante, iscrizione in dati_export_iscrizioni():
        mattino, ab_mattino = celle_export_fascia(partecipante, iscrizione, "mattino", laboratori)
        pomeriggio, ab_pomeriggio = celle_export_fascia(partecipante, iscrizione, "pomeriggio", laboratori)
        righe.append({
            "codice": partecipante.id,
            "nome": partecipante.nome,
            "cognome": partecipante.cognome,
            "email": partecipante.email,
            "sabato_mattina": mattino,
            "sottogruppo_mattino": ab_mattino,
            "sabato_pomeriggio": pomeriggio,
            "sottogruppo_pomeriggio": ab_pomeriggio,
            "domenica_mattina": descrivi_domenica_comunicazione(partecipante),
        })
    return righe


@app.route("/admin/comunicazioni")
@admin_required()
def genera_file_comunicazioni():
    return render_template("genera_comunicazioni.html")


@app.route("/admin/comunicazioni/valida", methods=["POST"])
@admin_required(api=True)
def valida_file_comunicazioni():
    return {"ok": False, "errore": "La validazione Excel non è più disponibile. Genera il file dai dati presenti nel sistema."}, 410


@app.route("/admin/comunicazioni/genera", methods=["POST"])
@admin_required(api=True)
def scarica_file_comunicazioni():
    workbook = None
    try:
        righe = leggi_dati_comunicazioni()
        workbook = crea_workbook_comunicazioni(righe)
        risposta = invia_file_excel(workbook, "comunicazioni_partecipanti")
        risposta.headers["X-Partecipanti-Esportati"] = str(len(righe))
        return risposta
    except Exception:
        app.logger.exception("Errore durante la generazione del file comunicazioni")
        return {"ok": False, "errore": "Non è stato possibile generare il file comunicazioni."}, 500
    finally:
        if workbook is not None:
            workbook.close()


@app.route("/admin/cambia_password", methods=["GET", "POST"])
@admin_required()
def cambia_password():
    if request.method == "POST":
        password_attuale = request.form.get("password_attuale", "")
        nuova_password = request.form.get("nuova_password", "")
        conferma_password = request.form.get("conferma_password", "")

        if not check_password_hash(current_user.password, password_attuale):
            flash("La password attuale non è corretta.", "warning")
        elif not nuova_password:
            flash("La nuova password non può essere vuota.", "warning")
        elif nuova_password != conferma_password:
            flash("La nuova password e la conferma non coincidono.", "warning")
        else:
            current_user.password = generate_password_hash(nuova_password)
            db.session.commit()
            flash("Password aggiornata correttamente.", "success")
            return redirect(url_for("cambia_password"))

    return render_template("cambia_password.html")

@app.route("/login", methods=["GET", "POST"])
def login():
    if request.method == "POST":
        utente = User.query.filter_by(username=request.form["username"]).first()
        if utente:
            if check_password_hash(utente.password, request.form["passwd"]):
                login_user(utente)
                return redirect(url_for("import_iscritti"))
            else:
                flash("Username o Password errati!", "warning")
        else:
            flash("Utente inesistente!", "warning")
    return render_template("login.html")

@app.route("/logout")
@login_required
def logout():
    logout_user()
    session.pop("temp_user", None)
    session.pop("identita_confermata", None)
    return redirect(url_for("index"))

@app.errorhandler(404)
def page_not_found(e):
    return render_template("errore_generico.html"), 404

@app.errorhandler(405)
def method_not_allowed(e):
    return render_template("errore_generico.html"), 405

@app.errorhandler(413)
def request_entity_too_large(e):
    return {
        "ok": False,
        "errore": "Il file supera la dimensione massima consentita di 10 MB.",
    }, 413

@app.errorhandler(500)
def internal_error(e):
    return render_template("errore_generico.html"), 500

if __name__ == "__main__":
    app.run(port=8000, host="0.0.0.0")
