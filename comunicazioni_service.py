from collections import Counter
import io
import os
import re
import zipfile

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from werkzeug.utils import secure_filename
from suddivisione_gruppi_service import NOMI_GRUPPI_DOMENICA


MAX_FILE_COMUNICAZIONI = 10 * 1024 * 1024
MAX_CONTENUTO_XLSX_ESTRATTO = 100 * 1024 * 1024
MAX_CODICE_CENSIMENTO = 2147483647
MIME_XLSX_CONSENTITI = {
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "application/vnd.ms-excel",
    "application/octet-stream",
    "application/zip",
    "application/x-zip-compressed",
}
COLONNE_SUDDIVISIONE_RICHIESTE = {
    "codice": "codice censimento",
    "nome": "nome",
    "cognome": "cognome",
    "email": "email",
    "gruppo": "gruppo",
}
COLONNE_OUTPUT_COMUNICAZIONI = (
    "Codice censimento",
    "Nome",
    "Cognome",
    "Email",
    "Sabato mattino",
    "Gruppo A/B mattino",
    "Sabato pomeriggio",
    "Gruppo A/B pomeriggio",
    "Domenica",
)


class ErroreComunicazioni(ValueError):
    def __init__(self, messaggio, duplicati=None):
        super().__init__(messaggio)
        self.duplicati = duplicati or []


def normalizza_intestazione(valore):
    if valore is None:
        return ""
    return " ".join(str(valore).strip().casefold().rstrip(":").split())


def valida_contenitore_xlsx(contenuto):
    try:
        with zipfile.ZipFile(io.BytesIO(contenuto)) as archivio:
            elementi = archivio.infolist()
            nomi = {elemento.filename for elemento in elementi}
            if "[Content_Types].xml" not in nomi or "xl/workbook.xml" not in nomi:
                raise ErroreComunicazioni("Il file non è un documento Excel .xlsx valido.")
            if len(elementi) > 2000:
                raise ErroreComunicazioni("Il file Excel contiene troppi elementi.")
            if sum(elemento.file_size for elemento in elementi) > MAX_CONTENUTO_XLSX_ESTRATTO:
                raise ErroreComunicazioni("Il contenuto del file Excel è troppo grande.")
    except zipfile.BadZipFile as errore:
        raise ErroreComunicazioni("Il file non è un documento Excel .xlsx valido.") from errore


def carica_excel_comunicazioni(file_caricato):
    if file_caricato is None or not file_caricato.filename:
        raise ErroreComunicazioni("Seleziona il file Excel della suddivisione domenicale.")

    nome_originale = file_caricato.filename
    if os.path.basename(nome_originale.replace("\\", "/")) != nome_originale:
        raise ErroreComunicazioni("Il nome del file non è valido.")
    nome_sicuro = secure_filename(nome_originale)
    if not nome_sicuro or not nome_sicuro.casefold().endswith(".xlsx"):
        raise ErroreComunicazioni("Sono accettati esclusivamente file .xlsx.")
    if file_caricato.mimetype and file_caricato.mimetype not in MIME_XLSX_CONSENTITI:
        raise ErroreComunicazioni("Il tipo del file caricato non è valido.")

    contenuto = file_caricato.stream.read(MAX_FILE_COMUNICAZIONI + 1)
    if not contenuto:
        raise ErroreComunicazioni("Il file caricato è vuoto.")
    if len(contenuto) > MAX_FILE_COMUNICAZIONI:
        raise ErroreComunicazioni("Il file supera la dimensione massima di 10 MB.")
    valida_contenitore_xlsx(contenuto)

    try:
        workbook = load_workbook(io.BytesIO(contenuto), data_only=True)
    except Exception as errore:
        raise ErroreComunicazioni("Il file Excel non è leggibile.") from errore
    return workbook


def normalizza_intero_positivo(valore, nome_campo, numero_riga):
    if isinstance(valore, bool):
        valore_normalizzato = None
    elif isinstance(valore, int):
        valore_normalizzato = valore
    elif isinstance(valore, float) and valore.is_integer():
        valore_normalizzato = int(valore)
    elif isinstance(valore, str) and re.fullmatch(r"[0-9]+", valore.strip()):
        valore_normalizzato = int(valore.strip())
    else:
        valore_normalizzato = None

    if (
        valore_normalizzato is None
        or valore_normalizzato <= 0
        or valore_normalizzato > MAX_CODICE_CENSIMENTO
    ):
        raise ErroreComunicazioni(
            f"La riga Excel {numero_riga} contiene un valore non valido per {nome_campo}."
        )
    return valore_normalizzato


def estrai_assegnazioni_domenica(workbook):
    if "Suddivisione" not in workbook.sheetnames:
        raise ErroreComunicazioni('Il file deve contenere il foglio "Suddivisione".')
    foglio = workbook["Suddivisione"]
    if foglio.max_row < 1:
        raise ErroreComunicazioni('Il foglio "Suddivisione" è vuoto.')

    colonne = {}
    for cella in foglio[1]:
        intestazione = normalizza_intestazione(cella.value)
        if intestazione and intestazione not in colonne:
            colonne[intestazione] = cella.column
    mancanti = [
        intestazione
        for intestazione in COLONNE_SUDDIVISIONE_RICHIESTE.values()
        if intestazione not in colonne
    ]
    if mancanti:
        raise ErroreComunicazioni(
            "Nel foglio Suddivisione mancano le colonne: " + ", ".join(mancanti) + "."
        )

    assegnazioni = {}
    codici_letti = []
    for numero_riga in range(2, foglio.max_row + 1):
        valori_riga = [cella.value for cella in foglio[numero_riga]]
        if all(valore is None or not str(valore).strip() for valore in valori_riga):
            continue
        codice = normalizza_intero_positivo(
            foglio.cell(numero_riga, colonne["codice censimento"]).value,
            "Codice censimento",
            numero_riga,
        )
        gruppo = normalizza_intero_positivo(
            foglio.cell(numero_riga, colonne["gruppo"]).value,
            "Gruppo",
            numero_riga,
        )
        codici_letti.append(codice)
        assegnazioni[codice] = {
            "codice": codice,
            "nome": foglio.cell(numero_riga, colonne["nome"]).value,
            "cognome": foglio.cell(numero_riga, colonne["cognome"]).value,
            "email": foglio.cell(numero_riga, colonne["email"]).value or "",
            "gruppo": gruppo,
        }

    duplicati = sorted(
        codice for codice, conteggio in Counter(codici_letti).items() if conteggio > 1
    )
    if duplicati:
        elenco = ", ".join(str(codice) for codice in duplicati)
        raise ErroreComunicazioni(
            f"Codici censimento duplicati nel file: {elenco}.",
            duplicati=duplicati,
        )
    if not assegnazioni:
        raise ErroreComunicazioni('Il foglio "Suddivisione" non contiene partecipanti.')
    return assegnazioni


def unisci_dati_comunicazioni(partecipanti_database, assegnazioni_domenica):
    codici_database = {partecipante["codice"] for partecipante in partecipanti_database}
    codici_excel = set(assegnazioni_domenica)
    anomalie = {
        "codici_excel_non_database": sorted(codici_excel - codici_database),
        "partecipanti_senza_domenica": sorted(codici_database - codici_excel),
        "iscrizioni_sabato_incomplete": sorted(
            partecipante["codice"]
            for partecipante in partecipanti_database
            if partecipante["sabato_incompleto"]
        ),
        "duplicati": [],
    }

    righe = []
    for partecipante in partecipanti_database:
        domenica = assegnazioni_domenica.get(partecipante["codice"])
        righe.append(
            {
                "codice": partecipante["codice"],
                "nome": partecipante["nome"],
                "cognome": partecipante["cognome"],
                "email": domenica["email"] if domenica else "",
                "sabato_mattina": partecipante["sabato_mattina"],
                "sabato_pomeriggio": partecipante["sabato_pomeriggio"],
                "domenica_mattina": domenica["gruppo"] if domenica else "",
            }
        )
    for codice in anomalie["codici_excel_non_database"]:
        domenica = assegnazioni_domenica[codice]
        righe.append(
            {
                "codice": domenica["codice"],
                "nome": domenica["nome"],
                "cognome": domenica["cognome"],
                "email": domenica["email"],
                "sabato_mattina": "",
                "sabato_pomeriggio": "",
                "domenica_mattina": domenica["gruppo"],
            }
        )
    return righe, anomalie


def descrivi_domenica_comunicazione(partecipante):
    if not partecipante.includi_domenica:
        return "Iscrizione non richiesta"
    if partecipante.gruppo_domenica is None:
        return "Non assegnato"
    return NOMI_GRUPPI_DOMENICA[partecipante.gruppo_domenica]


def crea_workbook_comunicazioni(righe):
    workbook = Workbook()
    foglio = workbook.active
    foglio.title = "Comunicazioni"
    foglio.append(COLONNE_OUTPUT_COMUNICAZIONI)
    for riga in righe:
        foglio.append(
            [
                riga["codice"],
                riga["nome"],
                riga["cognome"],
                riga["email"],
                riga["sabato_mattina"],
                riga["sottogruppo_mattino"],
                riga["sabato_pomeriggio"],
                riga["sottogruppo_pomeriggio"],
                riga["domenica_mattina"],
            ]
        )
    for cella in foglio[1]:
        cella.font = Font(bold=True, color="FFFFFF")
        cella.fill = PatternFill("solid", fgColor="0D6EFD")
    foglio.freeze_panes = "A2"
    foglio.auto_filter.ref = foglio.dimensions
    for indice, larghezza in enumerate((20, 24, 24, 34, 38, 22, 38, 24, 28), start=1):
        foglio.column_dimensions[foglio.cell(1, indice).column_letter].width = larghezza
    return workbook
