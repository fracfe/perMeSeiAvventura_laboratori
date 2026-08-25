from collections import Counter
import io
import os
import re
import zipfile

from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill
from werkzeug.utils import secure_filename

MAX_FILE_SUDDIVISIONE = 10 * 1024 * 1024
MAX_CONTENUTO_XLSX_ESTRATTO = 100 * 1024 * 1024
MAX_CODICE_CENSIMENTO = 2147483647
MIME_XLSX_CONSENTITI = {
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "application/vnd.ms-excel",
    "application/octet-stream",
    "application/zip",
    "application/x-zip-compressed",
}
CRITERI_SUDDIVISIONE = ("regione", "ruolo", "foca", "sesso")
INTESTAZIONI_SUDDIVISIONE = {
    "ruolo": "partecipo in qualità di",
    "foca": "foca",
    "sesso": "sesso",
}
INTESTAZIONI_IDENTIFICATIVE = {
    "codice": ("codice censimento", "codice socio", "codice"),
    "nome": ("nome",),
    "cognome": ("cognome",),
    "email": ("emailcontatto", "email", "e-mail"),
}


class ErroreSuddivisione(ValueError):
    pass


def normalizza_intestazione(valore):
    if valore is None:
        return ""
    return " ".join(str(valore).strip().casefold().rstrip(":").split())


def normalizza_valore_dinamico(valore, valori_canonici):
    if valore is None or not str(valore).strip():
        return "Non indicato"
    valore_pulito = " ".join(str(valore).strip().split())
    return valori_canonici.setdefault(valore_pulito.casefold(), valore_pulito)


def normalizza_foca(valore):
    valore_pulito = " ".join(str(valore or "").strip().casefold().split())
    return {
        "nomina": "Nomina",
        "cfa": "CFA",
        "cfm": "CFM",
    }.get(valore_pulito, "Sistema")


def normalizza_sesso(valore):
    valore_pulito = str(valore or "").strip().casefold()
    if valore_pulito == "m":
        return "M"
    if valore_pulito == "f":
        return "F"
    return "Non indicato"


def normalizza_codice_censimento(valore, numero_riga):
    if isinstance(valore, bool):
        codice = None
    elif isinstance(valore, int):
        codice = valore
    elif isinstance(valore, float) and valore.is_integer():
        codice = int(valore)
    elif isinstance(valore, str) and re.fullmatch(r"[0-9]+", valore.strip()):
        codice = int(valore.strip())
    else:
        codice = None

    if codice is None or codice <= 0 or codice > MAX_CODICE_CENSIMENTO:
        raise ErroreSuddivisione(
            f"La riga Excel {numero_riga} contiene un Codice censimento non valido: "
            f"deve essere un intero positivo non superiore a {MAX_CODICE_CENSIMENTO}."
        )
    return codice


def valida_contenitore_xlsx(contenuto):
    try:
        with zipfile.ZipFile(io.BytesIO(contenuto)) as archivio:
            elementi = archivio.infolist()
            nomi = {elemento.filename for elemento in elementi}
            if "[Content_Types].xml" not in nomi or "xl/workbook.xml" not in nomi:
                raise ErroreSuddivisione("Il file non è un documento Excel .xlsx valido.")
            if len(elementi) > 2000:
                raise ErroreSuddivisione("Il file Excel contiene troppi elementi.")
            if sum(elemento.file_size for elemento in elementi) > MAX_CONTENUTO_XLSX_ESTRATTO:
                raise ErroreSuddivisione("Il contenuto del file Excel è troppo grande.")
    except zipfile.BadZipFile as errore:
        raise ErroreSuddivisione("Il file non è un documento Excel .xlsx valido.") from errore


def carica_excel_suddivisione(file_caricato):
    if file_caricato is None or not file_caricato.filename:
        raise ErroreSuddivisione("Seleziona un file Excel .xlsx.")

    nome_originale = file_caricato.filename
    if os.path.basename(nome_originale.replace("\\", "/")) != nome_originale:
        raise ErroreSuddivisione("Il nome del file non è valido.")
    nome_sicuro = secure_filename(nome_originale)
    if not nome_sicuro or not nome_sicuro.casefold().endswith(".xlsx"):
        raise ErroreSuddivisione("Sono accettati esclusivamente file .xlsx.")
    if file_caricato.mimetype and file_caricato.mimetype not in MIME_XLSX_CONSENTITI:
        raise ErroreSuddivisione("Il tipo del file caricato non è valido.")

    contenuto = file_caricato.stream.read(MAX_FILE_SUDDIVISIONE + 1)
    if not contenuto:
        raise ErroreSuddivisione("Il file caricato è vuoto.")
    if len(contenuto) > MAX_FILE_SUDDIVISIONE:
        raise ErroreSuddivisione("Il file supera la dimensione massima di 10 MB.")
    valida_contenitore_xlsx(contenuto)

    try:
        workbook = load_workbook(io.BytesIO(contenuto), data_only=False)
    except Exception as errore:
        raise ErroreSuddivisione("Il file Excel non è leggibile.") from errore
    if not workbook.worksheets:
        raise ErroreSuddivisione("Il file Excel non contiene fogli.")
    return workbook


def trova_intestazioni_suddivisione(foglio):
    candidata_senza_regione_i = False
    for numero_riga in range(1, min(foglio.max_row, 50) + 1):
        colonne = {}
        for cella in foglio[numero_riga]:
            intestazione = normalizza_intestazione(cella.value)
            if intestazione and intestazione not in colonne:
                colonne[intestazione] = cella.column
        colonne_identificative = {
            campo: next(
                (colonne[alias] for alias in alias_consentiti if alias in colonne),
                None,
            )
            for campo, alias_consentiti in INTESTAZIONI_IDENTIFICATIVE.items()
        }
        if (
            all(nome in colonne for nome in INTESTAZIONI_SUDDIVISIONE.values())
            and all(colonne_identificative.values())
        ):
            if normalizza_intestazione(foglio.cell(numero_riga, 9).value) == "regione":
                colonne_criteri = {
                    criterio: colonne[intestazione]
                    for criterio, intestazione in INTESTAZIONI_SUDDIVISIONE.items()
                }
                return numero_riga, {**colonne_criteri, **colonne_identificative}
            candidata_senza_regione_i = True

    if candidata_senza_regione_i:
        raise ErroreSuddivisione(
            "La colonna I deve avere l'intestazione Regione nel file caricato."
        )
    raise ErroreSuddivisione(
        "Non trovo le intestazioni richieste: Codice censimento, Nome, Cognome, "
        "Email, Partecipo in qualità di:, FoCa, Sesso e Regione nella colonna I."
    )


def estrai_partecipanti_suddivisione(workbook):
    foglio = workbook.worksheets[0]
    riga_intestazioni, colonne = trova_intestazioni_suddivisione(foglio)
    canonici_regione = {}
    canonici_ruolo = {}
    partecipanti = []
    codici_letti = []
    for numero_riga in range(riga_intestazioni + 1, foglio.max_row + 1):
        if all(cella.value is None or str(cella.value).strip() == "" for cella in foglio[numero_riga]):
            continue
        identificativi = {
            campo: foglio.cell(numero_riga, colonne[campo]).value
            for campo in ("codice", "nome", "cognome", "email")
        }
        mancanti = [
            campo
            for campo, valore in identificativi.items()
            if valore is None or not str(valore).strip()
        ]
        if mancanti:
            raise ErroreSuddivisione(
                f"La riga Excel {numero_riga} non contiene Codice censimento, Nome, "
                "Cognome ed Email completi."
            )
        identificativi["codice"] = normalizza_codice_censimento(
            identificativi["codice"],
            numero_riga,
        )
        codici_letti.append(identificativi["codice"])
        partecipanti.append(
            {
                "indice": len(partecipanti),
                "riga_excel": numero_riga,
                **identificativi,
                "regione": normalizza_valore_dinamico(
                    foglio.cell(numero_riga, 9).value,
                    canonici_regione,
                ),
                "ruolo": normalizza_valore_dinamico(
                    foglio.cell(numero_riga, colonne["ruolo"]).value,
                    canonici_ruolo,
                ),
                "foca": normalizza_foca(
                    foglio.cell(numero_riga, colonne["foca"]).value
                ),
                "sesso": normalizza_sesso(
                    foglio.cell(numero_riga, colonne["sesso"]).value
                ),
            }
        )
    duplicati = sorted(
        codice for codice, conteggio in Counter(codici_letti).items() if conteggio > 1
    )
    if duplicati:
        elenco = ", ".join(str(codice) for codice in duplicati)
        raise ErroreSuddivisione(
            f"Codici censimento duplicati nel file: {elenco}."
        )
    if not partecipanti:
        raise ErroreSuddivisione("Il file non contiene partecipanti.")
    return foglio, riga_intestazioni, partecipanti


def valida_numero_gruppi(valore, numero_partecipanti):
    if isinstance(valore, bool) or not re.fullmatch(r"[0-9]+", str(valore or "")):
        raise ErroreSuddivisione("Il numero di laboratori deve essere un intero positivo.")
    numero_gruppi = int(valore)
    if numero_gruppi <= 0:
        raise ErroreSuddivisione("Il numero di laboratori deve essere maggiore di zero.")
    if numero_gruppi > numero_partecipanti:
        raise ErroreSuddivisione(
            "Il numero di laboratori non può superare il numero dei partecipanti."
        )
    return numero_gruppi


def punteggio_suddivisione(gruppi, partecipanti):
    """Restituisce un punteggio lessicografico nell'ordine dei criteri."""
    numero_gruppi = len(gruppi)
    risultato = []
    for criterio in CRITERI_SUDDIVISIONE:
        totali = Counter(partecipante[criterio] for partecipante in partecipanti)
        somma_quadrati = 0
        for gruppo in gruppi:
            conteggi = Counter(partecipanti[indice][criterio] for indice in gruppo)
            somma_quadrati += sum(conteggio * conteggio for conteggio in conteggi.values())
        risultato.append(
            numero_gruppi * somma_quadrati - sum(totale * totale for totale in totali.values())
        )
    return tuple(risultato)


def crea_suddivisione(partecipanti, numero_gruppi):
    numero_partecipanti = len(partecipanti)
    numero_gruppi = valida_numero_gruppi(numero_gruppi, numero_partecipanti)
    base, resto = divmod(numero_partecipanti, numero_gruppi)
    capacita = [base + (indice < resto) for indice in range(numero_gruppi)]
    frequenze = {
        criterio: Counter(partecipante[criterio] for partecipante in partecipanti)
        for criterio in CRITERI_SUDDIVISIONE
    }
    ordine = sorted(
        range(numero_partecipanti),
        key=lambda indice: (
            tuple(frequenze[criterio][partecipanti[indice][criterio]] for criterio in CRITERI_SUDDIVISIONE),
            indice,
        ),
    )

    gruppi = [[] for _ in range(numero_gruppi)]
    conteggi = [
        {criterio: Counter() for criterio in CRITERI_SUDDIVISIONE}
        for _ in range(numero_gruppi)
    ]
    for indice_partecipante in ordine:
        partecipante = partecipanti[indice_partecipante]
        disponibili = [
            indice for indice in range(numero_gruppi)
            if len(gruppi[indice]) < capacita[indice]
        ]
        gruppo_scelto = min(
            disponibili,
            key=lambda indice: (
                tuple(
                    conteggi[indice][criterio][partecipante[criterio]]
                    for criterio in CRITERI_SUDDIVISIONE
                ),
                len(gruppi[indice]) / capacita[indice],
                indice,
            ),
        )
        gruppi[gruppo_scelto].append(indice_partecipante)
        for criterio in CRITERI_SUDDIVISIONE:
            conteggi[gruppo_scelto][criterio][partecipante[criterio]] += 1

    punteggio = punteggio_suddivisione(gruppi, partecipanti)
    massimo_scambi = min(250, numero_partecipanti)
    for _ in range(massimo_scambi):
        scambio_migliore = None
        punteggio_migliore = punteggio
        for primo_gruppo in range(numero_gruppi - 1):
            for secondo_gruppo in range(primo_gruppo + 1, numero_gruppi):
                for posizione_primo, indice_primo in enumerate(gruppi[primo_gruppo]):
                    primo = partecipanti[indice_primo]
                    for posizione_secondo, indice_secondo in enumerate(gruppi[secondo_gruppo]):
                        secondo = partecipanti[indice_secondo]
                        if all(primo[c] == secondo[c] for c in CRITERI_SUDDIVISIONE):
                            continue
                        candidato = []
                        for posizione_criterio, criterio in enumerate(CRITERI_SUDDIVISIONE):
                            valore_primo = primo[criterio]
                            valore_secondo = secondo[criterio]
                            variazione = 0
                            if valore_primo != valore_secondo:
                                primo_a = conteggi[primo_gruppo][criterio][valore_primo]
                                primo_b = conteggi[primo_gruppo][criterio][valore_secondo]
                                secondo_a = conteggi[secondo_gruppo][criterio][valore_primo]
                                secondo_b = conteggi[secondo_gruppo][criterio][valore_secondo]
                                variazione = numero_gruppi * (
                                    (primo_a - 1) ** 2 - primo_a ** 2
                                    + (primo_b + 1) ** 2 - primo_b ** 2
                                    + (secondo_a + 1) ** 2 - secondo_a ** 2
                                    + (secondo_b - 1) ** 2 - secondo_b ** 2
                                )
                            candidato.append(punteggio[posizione_criterio] + variazione)
                        candidato = tuple(candidato)
                        if candidato < punteggio_migliore:
                            punteggio_migliore = candidato
                            scambio_migliore = (
                                primo_gruppo,
                                secondo_gruppo,
                                posizione_primo,
                                posizione_secondo,
                            )
        if scambio_migliore is None:
            break
        primo_gruppo, secondo_gruppo, posizione_primo, posizione_secondo = scambio_migliore
        indice_primo = gruppi[primo_gruppo][posizione_primo]
        indice_secondo = gruppi[secondo_gruppo][posizione_secondo]
        primo = partecipanti[indice_primo]
        secondo = partecipanti[indice_secondo]
        for criterio in CRITERI_SUDDIVISIONE:
            if primo[criterio] != secondo[criterio]:
                conteggi[primo_gruppo][criterio][primo[criterio]] -= 1
                conteggi[primo_gruppo][criterio][secondo[criterio]] += 1
                conteggi[secondo_gruppo][criterio][secondo[criterio]] -= 1
                conteggi[secondo_gruppo][criterio][primo[criterio]] += 1
        gruppi[primo_gruppo][posizione_primo], gruppi[secondo_gruppo][posizione_secondo] = (
            gruppi[secondo_gruppo][posizione_secondo],
            gruppi[primo_gruppo][posizione_primo],
        )
        punteggio = punteggio_migliore

    assegnazioni = [None] * numero_partecipanti
    for indice_gruppo, gruppo in enumerate(gruppi, start=1):
        for indice_partecipante in gruppo:
            assegnazioni[indice_partecipante] = indice_gruppo
    return assegnazioni, gruppi


def aggiungi_riepilogo_suddivisione(workbook, partecipanti, gruppi):
    if "Riepilogo" in workbook.sheetnames:
        workbook.remove(workbook["Riepilogo"])
    foglio = workbook.create_sheet("Riepilogo", 1)
    categorie = {
        criterio: sorted({p[criterio] for p in partecipanti}, key=str.casefold)
        for criterio in ("sesso", "foca", "regione", "ruolo")
    }
    intestazioni = ["Gruppo", "Totale partecipanti"]
    nomi_criteri = {
        "sesso": "Sesso",
        "foca": "FoCa",
        "regione": "Regione",
        "ruolo": "Partecipo in qualità di",
    }
    for criterio in ("sesso", "foca", "regione", "ruolo"):
        intestazioni.extend(
            f"{nomi_criteri[criterio]} — {categoria}"
            for categoria in categorie[criterio]
        )
    foglio.append(intestazioni)
    for cella in foglio[1]:
        cella.font = Font(bold=True, color="FFFFFF")
        cella.fill = PatternFill("solid", fgColor="198754")
    for indice_gruppo, gruppo in enumerate(gruppi, start=1):
        riga = [f"Gruppo {indice_gruppo}", len(gruppo)]
        for criterio in ("sesso", "foca", "regione", "ruolo"):
            conteggi = Counter(partecipanti[indice][criterio] for indice in gruppo)
            riga.extend(conteggi[categoria] for categoria in categorie[criterio])
        foglio.append(riga)
    foglio.freeze_panes = "C2"
    foglio.auto_filter.ref = foglio.dimensions
    foglio.column_dimensions["A"].width = 20
    foglio.column_dimensions["B"].width = 20
    for colonna in range(3, foglio.max_column + 1):
        foglio.column_dimensions[foglio.cell(1, colonna).column_letter].width = 28


def formatta_foglio_iscritti(foglio, larghezze):
    for cella in foglio[1]:
        cella.font = Font(bold=True, color="FFFFFF")
        cella.fill = PatternFill("solid", fgColor="198754")
    foglio.freeze_panes = "A2"
    foglio.auto_filter.ref = foglio.dimensions
    for indice, larghezza in enumerate(larghezze, start=1):
        foglio.column_dimensions[foglio.cell(1, indice).column_letter].width = larghezza


def genera_excel_suddivisione(workbook, partecipanti, numero_gruppi):
    assegnazioni, gruppi = crea_suddivisione(partecipanti, numero_gruppi)

    for foglio_originale in list(workbook.worksheets):
        workbook.remove(foglio_originale)
    foglio_generale = workbook.create_sheet("Suddivisione")
    foglio_generale.append(
        [
            "Codice censimento",
            "Nome",
            "Cognome",
            "Email",
            "Regione",
            "Partecipo in qualità di",
            "FoCa",
            "Sesso",
            "Gruppo",
        ]
    )
    for partecipante, assegnazione in zip(partecipanti, assegnazioni):
        foglio_generale.append(
            [
                partecipante["codice"],
                partecipante["nome"],
                partecipante["cognome"],
                partecipante["email"],
                partecipante["regione"],
                partecipante["ruolo"],
                partecipante["foca"],
                partecipante["sesso"],
                assegnazione,
            ]
        )
    formatta_foglio_iscritti(
        foglio_generale,
        (20, 24, 24, 34, 22, 32, 14, 16, 12),
    )

    aggiungi_riepilogo_suddivisione(workbook, partecipanti, gruppi)

    for numero_gruppo, gruppo in enumerate(gruppi, start=1):
        foglio_gruppo = workbook.create_sheet(f"Gruppo {numero_gruppo}")
        foglio_gruppo.append(["Codice censimento", "Nome", "Cognome", "Email"])
        for indice_partecipante in sorted(gruppo):
            partecipante = partecipanti[indice_partecipante]
            foglio_gruppo.append(
                [
                    partecipante["codice"],
                    partecipante["nome"],
                    partecipante["cognome"],
                    partecipante["email"],
                ]
            )
        formatta_foglio_iscritti(foglio_gruppo, (20, 24, 24, 34))

    return gruppi


def distribuzione_dimensioni(gruppi):
    return dict(Counter(len(gruppo) for gruppo in gruppi))
