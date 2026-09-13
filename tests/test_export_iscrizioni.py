import io
import os
import re
import unittest
from datetime import datetime
from unittest.mock import patch

os.environ["DB_TYPE"] = "sqlite"
os.environ["DB_NAME"] = ":memory:"
os.environ["SECRET_KEY"] = "test-secret-key"

from openpyxl import load_workbook
from werkzeug.security import generate_password_hash

from app import Iscrizione, Laboratorio, Partecipante, User, app, db


class ExportIscrizioniTestCase(unittest.TestCase):
    def setUp(self):
        app.config.update(TESTING=True)
        self.app_context = app.app_context()
        self.app_context.push()
        db.create_all()
        db.session.add(
            User(username="admin", password=generate_password_hash("password"))
        )
        db.session.commit()
        self.client = app.test_client()

    def tearDown(self):
        db.session.remove()
        db.drop_all()
        self.app_context.pop()

    def login_admin(self):
        return self.client.post(
            "/login",
            data={"username": "admin", "passwd": "password"},
        )

    def crea_dati(self):
        db.session.add_all(
            [
                Partecipante(id=101, nome="Mario", cognome="Rossi"),
                Partecipante(id=202, nome="Anna", cognome="Bianchi"),
                Partecipante(id=303, nome="Luca", cognome="Verdi"),
                Partecipante(id=404, nome="Non", cognome="Iniziato"),
                Partecipante(id=505, nome="Riga", cognome="Vuota"),
                Laboratorio(
                    id=1,
                    id_lab="M01",
                    titolo="Bosco",
                    descrizione="Laboratorio del mattino",
                    posti=20,
                    tipologia="mattino",
                ),
                Laboratorio(
                    id=2,
                    id_lab="P01",
                    titolo="Sentieri",
                    descrizione="Laboratorio del pomeriggio",
                    posti=20,
                    tipologia="pomeriggio",
                ),
                Laboratorio(
                    id=3,
                    id_lab="M02",
                    titolo="Tracce",
                    descrizione="Laboratorio senza iscritti",
                    posti=20,
                    tipologia="mattino",
                ),
            ]
        )
        db.session.commit()
        db.session.add_all(
            [
                Iscrizione(
                    data=datetime(2026, 8, 20, 10, 0),
                    partecipante=101,
                    scelta_mattino=1,
                    scelta_pomeriggio=2,
                ),
                Iscrizione(
                    data=datetime(2026, 8, 20, 10, 5),
                    partecipante=202,
                    scelta_mattino=1,
                    scelta_pomeriggio=None,
                ),
                Iscrizione(
                    data=datetime(2026, 8, 20, 10, 10),
                    partecipante=303,
                    scelta_mattino=None,
                    scelta_pomeriggio=2,
                ),
                Iscrizione(
                    data=datetime(2026, 8, 20, 10, 15),
                    partecipante=505,
                    scelta_mattino=None,
                    scelta_pomeriggio=None,
                ),
            ]
        )
        db.session.commit()

    intestazioni = ("Codice censimento", "Nome", "Cognome", "Gruppo", "Zona", "Regione", "FoCa", "Email")
    route = "/admin/iscrizioni/esporta/laboratori"

    def leggi_workbook(self, risposta):
        self.assertEqual(risposta.status_code, 200)
        self.assertEqual(risposta.mimetype, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        self.assertRegex(risposta.headers["Content-Disposition"],
                         r"attachment; filename=iscrizioni_per_laboratorio_\d{4}-\d{2}-\d{2}\.xlsx")
        workbook = load_workbook(io.BytesIO(risposta.data))
        self.addCleanup(workbook.close)
        for foglio in workbook:
            extra = (("Laboratorio mattino", "Gruppo A/B mattino", "Laboratorio pomeriggio", "Gruppo A/B pomeriggio")
                     if foglio.title == "Tutti i partecipanti" else
                     () if "Non partecipa" in foglio.title else ("Gruppo A/B",))
            self.assertEqual(next(foglio.values), self.intestazioni + extra)
            self.assertEqual(foglio.freeze_panes, "A2")
            self.assertTrue(foglio.auto_filter.ref)
        return workbook

    def codici(self, workbook, nome):
        return [r[0] for r in workbook[nome].iter_rows(min_row=2, values_only=True)]

    def test_workbook_unico_anagrafica_completa_e_fogli_laboratori(self):
        self.crea_dati()
        p = db.session.get(Partecipante, 101)
        p.gruppo, p.zona, p.regione, p.foca, p.email = "Roma 1", "Roma", "Lazio", "FoCa originale", "mario@example.test"
        db.session.commit()
        self.login_admin()
        wb = self.leggi_workbook(self.client.get(self.route))
        self.assertEqual(wb.sheetnames, ["Tutti i partecipanti", "M - M01 - Bosco", "M - M02 - Tracce",
                                        "P - P01 - Sentieri", "M - Non partecipa", "P - Non partecipa"])
        self.assertEqual(self.codici(wb, "Tutti i partecipanti"), [202, 404, 101, 303, 505])
        self.assertEqual(self.codici(wb, "M - M01 - Bosco"), [202, 101])
        self.assertEqual(self.codici(wb, "P - P01 - Sentieri"), [101, 303])
        self.assertEqual(self.codici(wb, "M - M02 - Tracce"), [])
        attesa = (101, "Mario", "Rossi", "Roma 1", "Roma", "Lazio", "FoCa originale", "mario@example.test")
        for nome in ("Tutti i partecipanti", "M - M01 - Bosco", "P - P01 - Sentieri"):
            self.assertIn(attesa, [r[:8] for r in list(wb[nome].values)[1:]])

    def test_rinunce_esclusi_e_nessuna_modifica_alle_assegnazioni(self):
        self.crea_dati()
        db.session.get(Partecipante, 202).deve_iscriversi_sabato = False
        i = Iscrizione.query.filter_by(partecipante=202).one()
        i.non_partecipa_pomeriggio = True
        i.sottogruppo_mattino = "A"
        i = Iscrizione.query.filter_by(partecipante=101).one()
        i.sottogruppo_mattino, i.sottogruppo_pomeriggio = "B", "A"
        Iscrizione.query.filter_by(partecipante=303).one().non_partecipa_mattino = True
        Iscrizione.query.filter_by(partecipante=505).one().non_partecipa_pomeriggio = True
        db.session.get(Partecipante, 101).gruppo_domenica = 7
        db.session.commit()
        self.login_admin()
        def stato():
            return {m.__tablename__: [tuple(getattr(r, c.name) for c in m.__table__.columns)
                    for r in m.query.order_by(m.id)] for m in (Partecipante, Iscrizione, Laboratorio)}
        prima = stato()
        with patch("app.ricalcola_sottogruppi_ab", side_effect=AssertionError("non ricalcolare")), \
             patch("app.calcola_sottogruppi_ab", side_effect=AssertionError("non ricalcolare")), \
             patch.object(db.session, "commit", side_effect=AssertionError("non scrivere")):
            wb = self.leggi_workbook(self.client.get(self.route))
        self.assertEqual(self.codici(wb, "Tutti i partecipanti"), [202, 404, 101, 303, 505])
        self.assertEqual(self.codici(wb, "M - M01 - Bosco"), [101])
        self.assertEqual(self.codici(wb, "M - Non partecipa"), [303])
        self.assertEqual(self.codici(wb, "P - Non partecipa"), [505])
        righe = {r[0]: r for r in wb["Tutti i partecipanti"].iter_rows(min_row=2, values_only=True)}
        self.assertEqual(righe[202][8:], ("Iscrizione non richiesta", None, "Iscrizione non richiesta", None))
        self.assertEqual(righe[404][8:], ("Non iscritto", None, "Non iscritto", None))
        self.assertEqual(righe[101][8:], ("M01 - Bosco", "B", "P01 - Sentieri", "A"))
        self.assertEqual(righe[303][8:], ("Non partecipa", None, "P01 - Sentieri", None))
        self.assertEqual(righe[505][8:], ("Non iscritto", None, "Non partecipa", None))
        self.assertEqual(list(wb["M - M01 - Bosco"].values)[1][-1], "B")
        self.assertEqual([r[-1] for r in list(wb["P - P01 - Sentieri"].values)[1:]], ["A", None])
        self.assertEqual(stato(), prima)

    def test_ordinamento_cognome_nome_codice_in_tutti_i_fogli(self):
        self.crea_dati()
        for codice, nome, cognome in ((909, "Anna", "Rossi"), (808, "Anna", "Rossi"), (707, "Zeno", "Rossi")):
            db.session.add(Partecipante(id=codice, nome=nome, cognome=cognome))
            db.session.flush()
            db.session.add(Iscrizione(partecipante=codice, scelta_mattino=1,
                                      non_partecipa_pomeriggio=True, data=datetime(2026, 9, 13)))
        db.session.commit()
        self.login_admin()
        wb = self.leggi_workbook(self.client.get(self.route))
        for foglio in wb:
            righe = list(foglio.values)[1:]
            self.assertEqual(righe, sorted(righe, key=lambda r: (r[2], r[1], r[0])))
        self.assertEqual(self.codici(wb, "M - Non partecipa"), [])
        self.assertEqual(self.codici(wb, "P - Non partecipa"), [808, 909, 707])

    def test_export_valido_con_zero_iscrizioni_o_nessun_partecipante(self):
        self.login_admin()
        wb = self.leggi_workbook(self.client.get(self.route))
        self.assertEqual(len(wb.sheetnames), 3)
        self.assertTrue(all(f.max_row == 1 for f in wb))
        db.session.add_all([Partecipante(id=1, nome="Uno", cognome="Rossi"),
                            Partecipante(id=2, nome="Due", cognome="Verdi", deve_iscriversi_sabato=False)])
        db.session.commit()
        wb = self.leggi_workbook(self.client.get(self.route))
        self.assertEqual(self.codici(wb, "Tutti i partecipanti"), [1, 2])
        self.assertEqual(Iscrizione.query.count(), 0)

    def test_navigazione_e_rimozione_vecchi_export(self):
        self.login_admin()
        pagina = self.client.get("/admin/iscrizioni").get_data(as_text=True)
        self.assertIn('href="' + self.route + '">Esporta iscrizioni sabato</a>', pagina)
        self.assertIn('>Suddivisione gruppi domenica</a>', pagina)
        self.assertNotIn('>Suddivisione gruppi</a>', pagina)
        self.assertNotIn('>Esporta iscrizioni</a>', pagina)
        for route in ("/admin/iscrizioni/esporta", "/admin/iscrizioni/esporta/elenco"):
            self.assertNotIn('href="' + route + '"', pagina)
            self.assertEqual(self.client.get(route).status_code, 404)

    def test_export_non_disponibile_senza_autenticazione_admin(self):
        endpoint = (
            "/admin/iscrizioni/esporta/laboratori",
        )
        for percorso in endpoint:
            with self.subTest(percorso=percorso, utente="anonimo"):
                self.assertEqual(self.client.get(percorso).status_code, 302)

        db.session.add(Partecipante(id=999, nome="Utente", cognome="Temporaneo"))
        db.session.commit()
        self.client.post("/verifica_iscrizione", json={"codice_socio": 999})

        for percorso in endpoint:
            with self.subTest(percorso=percorso, utente="partecipante"):
                partecipante = self.client.get(percorso)
                self.assertEqual(partecipante.status_code, 302)
                self.assertEqual(partecipante.headers["Location"], "/")

        self.client.get("/logout")
        db.session.add(
            User(
                username="operatore",
                password=generate_password_hash("password"),
            )
        )
        db.session.commit()
        self.client.post(
            "/login",
            data={"username": "operatore", "passwd": "password"},
        )
        for percorso in endpoint:
            with self.subTest(percorso=percorso, utente="non_admin"):
                operatore = self.client.get(percorso)
                self.assertEqual(operatore.status_code, 302)
                self.assertEqual(operatore.headers["Location"], "/")

    def test_nomi_foglio_lunghi_non_validi_e_duplicati_restano_univoci(self):
        db.session.add_all(
            [
                Laboratorio(
                    id=10,
                    id_lab="DUP",
                    titolo="Titolo/estremamente lungo che continua primo",
                    descrizione="Primo",
                    posti=10,
                    tipologia="mattino",
                ),
                Laboratorio(
                    id=11,
                    id_lab="DUP",
                    titolo="Titolo?estremamente lungo che continua secondo",
                    descrizione="Secondo",
                    posti=10,
                    tipologia="mattino",
                ),
            ]
        )
        db.session.commit()
        self.login_admin()

        risposta = self.client.get("/admin/iscrizioni/esporta/laboratori")
        workbook = self.leggi_workbook(risposta)

        nomi_duplicati = [
            nome for nome in workbook.sheetnames if nome.startswith("M - DUP")
        ]
        self.assertEqual(len(workbook.sheetnames), 5)
        self.assertEqual(len(nomi_duplicati), 2)
        self.assertEqual(len({nome.casefold() for nome in workbook.sheetnames}), 5)
        self.assertTrue(any(nome.endswith(" (2)") for nome in nomi_duplicati))
        for nome in workbook.sheetnames:
            self.assertLessEqual(len(nome), 31)
            self.assertIsNone(re.search(r"[\\/*?:\[\]]", nome))
