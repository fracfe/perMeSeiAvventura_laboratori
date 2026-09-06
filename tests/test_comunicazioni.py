import io
import os
import unittest
from contextlib import ExitStack
from datetime import datetime
from unittest.mock import patch

os.environ["DB_TYPE"] = "sqlite"
os.environ["DB_NAME"] = ":memory:"
os.environ["SECRET_KEY"] = "test-secret-key"

from openpyxl import load_workbook
from werkzeug.security import generate_password_hash
from app import Iscrizione, Laboratorio, Partecipante, SysOption, User, app, db


class ComunicazioniTestCase(unittest.TestCase):
    def setUp(self):
        app.config.update(TESTING=True)
        self.context = app.app_context()
        self.context.push()
        db.create_all()
        db.session.add(User(username="admin", password=generate_password_hash("password")))
        db.session.commit()
        self.client = app.test_client()

    def tearDown(self):
        db.session.remove()
        db.drop_all()
        self.context.pop()

    def login_admin(self):
        self.client.post("/login", data={"username": "admin", "passwd": "password"})

    def crea_dati(self):
        db.session.add_all([
            Partecipante(id=i, nome="Mario", cognome="Rossi", email=f"persona{i}@example.test",
                        includi_domenica=True, gruppo_domenica=gruppo)
            for i, gruppo in ((1, 1), (2, 7), (3, 20), (4, None), (5, 7))
        ])
        db.session.add_all([
            Laboratorio(id=1, id_lab="M01", titolo="Bosco", descrizione="", posti=20, tipologia="mattino"),
            Laboratorio(id=2, id_lab="P01", titolo="Sentieri", descrizione="", posti=20, tipologia="pomeriggio"),
        ])
        db.session.commit()
        db.session.get(Partecipante, 5).deve_iscriversi_sabato = False
        db.session.get(Partecipante, 5).includi_domenica = False
        db.session.add_all([
            Iscrizione(partecipante=1, scelta_mattino=1, scelta_pomeriggio=2,
                       sottogruppo_mattino="B", sottogruppo_pomeriggio="A", data=datetime(2026, 8, 20)),
            Iscrizione(partecipante=2, non_partecipa_mattino=True, scelta_pomeriggio=2,
                       sottogruppo_mattino="A", data=datetime(2026, 8, 20)),
            Iscrizione(partecipante=3, scelta_mattino=1, non_partecipa_pomeriggio=True,
                       sottogruppo_pomeriggio="B", data=datetime(2026, 8, 20)),
            Iscrizione(partecipante=5, scelta_mattino=1, non_partecipa_pomeriggio=True,
                       sottogruppo_mattino="A", sottogruppo_pomeriggio="B", data=datetime(2026, 8, 20)),
        ])
        db.session.commit()

    def scarica(self):
        risposta = self.client.post("/admin/comunicazioni/genera")
        self.assertEqual(risposta.status_code, 200)
        self.assertEqual(risposta.content_type, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        self.assertRegex(risposta.headers["Content-Disposition"], r"comunicazioni_partecipanti_\d{4}-\d{2}-\d{2}\.xlsx")
        workbook = load_workbook(io.BytesIO(risposta.data), read_only=True)
        self.addCleanup(workbook.close)
        self.assertEqual(workbook.sheetnames, ["Comunicazioni"])
        righe = list(workbook.active.iter_rows(values_only=True))
        self.assertEqual(int(risposta.headers["X-Partecipanti-Esportati"]), len(righe) - 1)
        self.assertNotIn("X-Anomalie-Totali", risposta.headers)
        return righe

    def test_colonne_tutti_partecipanti_email_stati_e_ab_persistiti(self):
        self.crea_dati()
        self.login_admin()
        righe = self.scarica()
        self.assertEqual(righe[0], ("Codice censimento", "Nome", "Cognome", "Email", "Sabato mattino",
                                    "Gruppo A/B mattino", "Sabato pomeriggio", "Gruppo A/B pomeriggio", "Domenica"))
        self.assertEqual(len(righe), 6)
        self.assertEqual([r[0] for r in righe[1:]], [1, 2, 3, 4, 5])
        per_codice = {r[0]: r for r in righe[1:]}
        self.assertEqual(per_codice[1], (1, "Mario", "Rossi", "persona1@example.test", "M01 - Bosco", "B", "P01 - Sentieri", "A", "Avventura"))
        self.assertEqual(per_codice[2][4:], ("Non partecipa", "Non partecipa", "P01 - Sentieri", "Non assegnato", "Gradualità"))
        self.assertEqual(per_codice[3][4:], ("M01 - Bosco", "Non assegnato", "Non partecipa", "Non partecipa", "Vivere"))
        self.assertEqual(per_codice[4][4:], ("Non iscritto", "Non iscritto", "Non iscritto", "Non iscritto", "Non assegnato"))
        self.assertEqual(per_codice[5][4:], ("Iscrizione non richiesta", "Iscrizione non richiesta", "Iscrizione non richiesta", "Iscrizione non richiesta", "Iscrizione non richiesta"))
        for i in range(1, 6):
            self.assertEqual(per_codice[i][3], f"persona{i}@example.test")

    def test_generazione_senza_upload_merge_ricalcoli_o_modifiche(self):
        self.crea_dati()
        self.login_admin()
        def stato():
            db.session.expire_all()
            return {m.__tablename__: [tuple(getattr(r, c.name) for c in m.__table__.columns)
                    for r in m.query.order_by(*m.__table__.primary_key.columns)]
                    for m in (Partecipante, Iscrizione, Laboratorio, User, SysOption)}
        prima = stato()
        with ExitStack() as stack:
            for funzione in ("comunicazioni_service.carica_excel_comunicazioni",
                             "comunicazioni_service.estrai_assegnazioni_domenica",
                             "comunicazioni_service.unisci_dati_comunicazioni",
                             "app.ricalcola_gruppi_domenica", "app.calcola_gruppi_domenica",
                             "app.ricalcola_sottogruppi_ab", "app.calcola_sottogruppi_ab"):
                stack.enter_context(patch(funzione, side_effect=AssertionError("non utilizzare")))
            self.scarica()
        self.assertEqual(stato(), prima)

    def test_database_vuoto_e_senza_iscrizioni_email_mancante(self):
        self.login_admin()
        self.assertEqual(len(self.scarica()), 1)
        db.session.add_all([
            Partecipante(id=1, nome="Anna", cognome="Rossi"),
            Partecipante(id=2, nome="Luca", cognome="Verdi", deve_iscriversi_sabato=False, includi_domenica=True),
        ])
        db.session.commit()
        righe = self.scarica()
        self.assertEqual(righe[1], (1, "Anna", "Rossi", None, "Non iscritto", "Non iscritto", "Non iscritto", "Non iscritto", "Iscrizione non richiesta"))
        self.assertEqual(righe[2][4:], ("Iscrizione non richiesta", "Iscrizione non richiesta", "Iscrizione non richiesta", "Iscrizione non richiesta", "Non assegnato"))
        self.assertEqual(Iscrizione.query.count(), 0)

    def test_ordinamento_cognome_nome_codice_deterministico(self):
        db.session.add_all([
            Partecipante(id=9, nome="Anna", cognome="Rossi"),
            Partecipante(id=2, nome="Anna", cognome="Rossi"),
            Partecipante(id=3, nome="Zeno", cognome="Bianchi"),
            Partecipante(id=1, nome="Luca", cognome="Rossi"),
        ])
        db.session.commit()
        self.login_admin()
        prima = self.scarica()
        self.assertEqual([r[0] for r in prima[1:]], [3, 2, 9, 1])
        self.assertEqual(self.scarica(), prima)

    def test_pagina_senza_vecchio_flusso_e_validazione_dismessa(self):
        self.login_admin()
        pagina = self.client.get("/admin/comunicazioni")
        self.assertEqual(pagina.status_code, 200)
        self.assertIn(b"Genera file comunicazioni", pagina.data)
        self.assertIn(b'action="/admin/comunicazioni/genera"', pagina.data)
        self.assertIn(b'aria-current="page"', pagina.data)
        for testo in (b'type="file"', b"/admin/comunicazioni/valida", b"EmailReferente", b"merge", b"solo Excel"):
            self.assertNotIn(testo, pagina.data)
        self.assertEqual(self.client.post("/admin/comunicazioni/valida").status_code, 410)
        self.assertEqual(self.client.get("/admin/comunicazioni/genera").status_code, 405)

    def test_route_protette_admin(self):
        def verifica(anonimo=False):
            self.assertEqual(self.client.get("/admin/comunicazioni").status_code, 302)
            self.assertEqual(self.client.post("/admin/comunicazioni/genera").status_code, 302 if anonimo else 403)
            self.assertEqual(self.client.post("/admin/comunicazioni/valida").status_code, 302 if anonimo else 403)
        verifica(anonimo=True)
        db.session.add(Partecipante(id=1, nome="Mario", cognome="Rossi"))
        db.session.commit()
        self.client.post("/verifica_iscrizione", json={"codice_socio": 1})
        verifica()
        db.session.add(User(username="operatore", password=generate_password_hash("password")))
        db.session.commit()
        self.client.post("/login", data={"username": "operatore", "passwd": "password"})
        verifica()

    def test_errori_interni_restituiscono_500_senza_dettagli(self):
        self.login_admin()
        for funzione in ("app.leggi_dati_comunicazioni", "app.crea_workbook_comunicazioni"):
            with self.subTest(funzione=funzione), patch(funzione, side_effect=RuntimeError("simulato")):
                risposta = self.client.post("/admin/comunicazioni/genera")
                self.assertEqual(risposta.status_code, 500)
                self.assertFalse(risposta.get_json()["ok"])
                self.assertNotIn("simulato", risposta.get_json()["errore"])
