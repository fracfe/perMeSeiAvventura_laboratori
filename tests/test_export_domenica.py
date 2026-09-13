import io
import os
import unittest
from contextlib import ExitStack
from datetime import datetime
from unittest.mock import patch

os.environ['DB_TYPE'] = 'sqlite'
os.environ['DB_NAME'] = ':memory:'
os.environ['SECRET_KEY'] = 'test-secret-key'

from openpyxl import load_workbook
from werkzeug.security import generate_password_hash
from app import app, db, Partecipante, User, Iscrizione, SysOption, NOMI_GRUPPI_DOMENICA


class ExportDomenicaTestCase(unittest.TestCase):
    route = '/admin/suddivisione-gruppi/esporta'
    intestazioni = ('Codice censimento', 'Nome', 'Cognome', 'Gruppo', 'Zona', 'Regione', 'FoCa', 'Email')

    def setUp(self):
        app.config.update(TESTING=True)
        self.context = app.app_context()
        self.context.push()
        db.create_all()
        db.session.add(User(username='admin', password=generate_password_hash('password')))
        db.session.commit()
        self.client = app.test_client()

    def tearDown(self):
        db.session.remove()
        db.drop_all()
        self.context.pop()

    def login(self):
        self.client.post('/login', data={'username': 'admin', 'passwd': 'password'})

    def dati(self):
        db.session.add_all([
            Partecipante(id=9, nome='Zeno', cognome='Rossi', foca='Valore FoCa originale', gruppo='Roma 1', zona='Roma', email='zeno@example.test', regione='Lazio', includi_domenica=True, gruppo_domenica=1),
            Partecipante(id=7, nome='Anna', cognome='Rossi', foca='CFM', includi_domenica=True, gruppo_domenica=1),
            Partecipante(id=3, nome='Anna', cognome='Rossi', foca='CFA', includi_domenica=True, gruppo_domenica=1),
            Partecipante(id=5, nome='Marta', cognome='Bianchi', includi_domenica=True, gruppo_domenica=7),
            Partecipante(id=2, nome='Luca', cognome='Verdi', includi_domenica=True, gruppo_domenica=20),
            Partecipante(id=4, nome='Non', cognome='Assegnato', includi_domenica=True),
            Partecipante(id=8, nome='Escluso', cognome='Rossi', includi_domenica=False, gruppo_domenica=1),
            Partecipante(id=6, nome='Escluso', cognome='Bianchi', includi_domenica=False),
            SysOption(key='ultimo_ricalcolo_domenica', value='prima'),
        ])
        db.session.flush()
        db.session.add(Iscrizione(partecipante=9, data=datetime(2026, 9, 1), non_partecipa_mattino=True))
        db.session.commit()

    def scarica(self):
        risposta = self.client.get(self.route)
        self.assertEqual(risposta.status_code, 200)
        self.assertEqual(risposta.mimetype, 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        self.assertRegex(risposta.headers['Content-Disposition'], r'attachment; filename=gruppi_domenica_\d{4}-\d{2}-\d{2}\.xlsx')
        wb = load_workbook(io.BytesIO(risposta.data))
        self.addCleanup(wb.close)
        return wb

    def test_richiede_admin_anche_per_utente_partecipante(self):
        self.assertEqual(self.client.get(self.route).status_code, 302)
        self.dati()
        self.client.post('/verifica_iscrizione', json={'codice_socio': 9})
        self.client.post('/conferma_identita')
        self.assertNotEqual(self.client.get(self.route).status_code, 200)
        self.login()
        self.scarica()

    def test_fogli_intestazioni_inclusi_nomi_ordinamento_e_foca(self):
        self.dati()
        self.login()
        wb = self.scarica()
        self.assertEqual(wb.sheetnames, ['Tutti i partecipanti'] + list(NOMI_GRUPPI_DOMENICA.values()))
        generale = list(wb.worksheets[0].values)
        self.assertEqual(generale[0], self.intestazioni + ("Gruppo domenica",))
        self.assertEqual([r[0] for r in generale[1:]], [4, 5, 3, 7, 9, 2])
        righe = {r[0]: r for r in generale[1:]}
        self.assertEqual(righe[4][-1], "Non assegnato")
        self.assertEqual(righe[5][-1], "Gradualità")
        self.assertEqual(righe[2][-1], "Vivere")
        self.assertEqual(righe[9], (9, 'Zeno', 'Rossi', 'Roma 1', 'Roma', 'Lazio', 'Valore FoCa originale', 'zeno@example.test', 'Avventura'))
        self.assertIsNone(righe[5][6])
        gruppi_attesi = {'Avventura': [3, 7, 9], 'Gradualità': [5], 'Vivere': [2]}
        for nome in NOMI_GRUPPI_DOMENICA.values():
            foglio = wb[nome]
            self.assertEqual(next(foglio.values), self.intestazioni)
            self.assertEqual(list(foglio.values)[1:], [righe[c][:8] for c in gruppi_attesi.get(nome, [])])
            self.assertEqual(foglio.freeze_panes, 'A2')
            self.assertTrue(foglio.auto_filter.ref)

    def test_export_senza_ricalcoli_o_scritture(self):
        self.dati()
        self.login()
        def snapshot():
            return {m.__tablename__: [tuple(getattr(r, c.name) for c in m.__table__.columns)
                    for r in m.query.order_by(*m.__table__.primary_key.columns)]
                    for m in (Partecipante, Iscrizione, SysOption)}
        prima = snapshot()
        with ExitStack() as stack:
            for funzione in ('ricalcola_gruppi_domenica', 'calcola_gruppi_domenica', 'ricalcola_sottogruppi_ab'):
                stack.enter_context(patch('app.' + funzione, side_effect=AssertionError('Non ricalcolare')))
            stack.enter_context(patch.object(db.session, 'commit', side_effect=AssertionError('Non scrivere')))
            self.scarica()
        self.assertEqual(snapshot(), prima)

    def test_zero_partecipanti_fogli_vuoti_validi_e_pulsante(self):
        self.login()
        wb = self.scarica()
        self.assertEqual(len(wb.sheetnames), 21)
        self.assertTrue(all(f.max_row == 1 for f in wb))
        pagina = self.client.get('/admin/suddivisione-gruppi').get_data(as_text=True)
        self.assertIn('Esporta gruppi domenica', pagina)
        self.assertIn('href="' + self.route + '"', pagina)
