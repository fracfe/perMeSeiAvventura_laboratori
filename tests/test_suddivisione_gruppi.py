import io
import os
import unittest
from unittest.mock import patch

os.environ["DB_TYPE"] = "sqlite"
os.environ["DB_NAME"] = ":memory:"
os.environ["SECRET_KEY"] = "test-secret-key"

from openpyxl import Workbook, load_workbook
from werkzeug.security import generate_password_hash

from app import Iscrizione, Laboratorio, Partecipante, SysOption, User, app, db
from suddivisione_gruppi_service import (
    ErroreSuddivisione,
    crea_suddivisione,
    estrai_partecipanti_suddivisione,
    genera_excel_suddivisione,
    normalizza_foca,
    normalizza_sesso,
    punteggio_suddivisione,
)


def partecipante(indice, regione="Lazio", ruolo="Capo", foca="Sistema", sesso="M"):
    return {
        "indice": indice,
        "riga_excel": indice + 2,
        "codice": indice + 1000,
        "nome": f"Nome {indice}",
        "cognome": f"Cognome {indice}",
        "email": f"persona{indice}@example.test",
        "regione": regione,
        "ruolo": ruolo,
        "foca": foca,
        "sesso": sesso,
    }


def crea_workbook_excel(numero_partecipanti=12):
    workbook = Workbook()
    foglio = workbook.active
    foglio.title = "Iscritti"
    foglio.append(["Titolo report"])
    foglio.append(
        [
            "Partecipo in qualità di:",
            "FoCa",
            "Sesso",
            "Nome",
            "Cognome",
            "Altro 1",
            "Altro 2",
            "Regione",
            "Regione",
            "Codice censimento",
            "Email",
        ]
    )
    for indice in range(numero_partecipanti):
        foglio.append(
            [
                "Categoria rara" if indice < 3 else f"Ruolo {indice % 2}",
                [" Nomina ", "cfa", "CFM", "altro"][indice % 4],
                ["M", "f", "", "X"][indice % 4],
                f"Nome {indice}",
                f"Cognome {indice}",
                None,
                None,
                "Regione sbagliata",
                f"Regione {indice % 3}",
                indice + 1000,
                f"persona{indice}@example.test",
            ]
        )
    foglio.append([None] * 10)
    return workbook


def crea_workbook_excel_intestazioni_reali():
    workbook = Workbook()
    foglio = workbook.active
    foglio.title = "Iscritti"
    intestazioni = {
        2: "Regione",
        3: "Codice",
        4: "Cognome",
        5: "Nome",
        7: "Sesso",
        8: "FoCa",
        9: "Regione",
        20: "EmailContatto",
        22: "Partecipo in qualità di:",
    }
    for colonna, intestazione in intestazioni.items():
        foglio.cell(1, colonna, intestazione)
    valori = {
        2: "Regione da ignorare",
        3: 123,
        4: "Rossi",
        5: "Mario",
        7: "M",
        8: "CFA",
        9: "Lazio",
        20: "mario@example.test",
        22: "Capo",
    }
    for colonna, valore in valori.items():
        foglio.cell(2, colonna, valore)
    return workbook


def workbook_bytes(workbook):
    contenuto = io.BytesIO()
    workbook.save(contenuto)
    contenuto.seek(0)
    return contenuto


class AlgoritmoSuddivisioneTestCase(unittest.TestCase):
    def test_integrita_dimensioni_e_casi_non_divisibili(self):
        for totale, gruppi_attesi in ((1, 1), (10, 3), (20, 10), (41, 6), (61, 4)):
            with self.subTest(totale=totale, gruppi=gruppi_attesi):
                dati = [
                    partecipante(
                        indice,
                        regione=f"R{indice % 5}",
                        ruolo=f"Q{indice % 4}",
                        foca=["Nomina", "CFA", "CFM", "Sistema"][indice % 4],
                        sesso="M" if indice % 2 else "F",
                    )
                    for indice in range(totale)
                ]
                assegnazioni, gruppi = crea_suddivisione(dati, gruppi_attesi)
                assegnati = [indice for gruppo in gruppi for indice in gruppo]
                dimensioni = [len(gruppo) for gruppo in gruppi]

                self.assertEqual(len(assegnati), totale)
                self.assertEqual(sorted(assegnati), list(range(totale)))
                self.assertEqual(len(set(assegnati)), totale)
                self.assertEqual(len(assegnazioni), totale)
                self.assertNotIn(None, assegnazioni)
                self.assertLessEqual(max(dimensioni) - min(dimensioni), 1)
                self.assertTrue(
                    all(
                        dimensione in (totale // gruppi_attesi, -(-totale // gruppi_attesi))
                        for dimensione in dimensioni
                    )
                )

    def test_regione_uniforme_e_categoria_rara_sono_distribuite(self):
        dati_regione = [partecipante(indice, regione="Piemonte") for indice in range(20)]
        _, gruppi_regione = crea_suddivisione(dati_regione, 10)
        self.assertEqual([len(gruppo) for gruppo in gruppi_regione], [2] * 10)

        dati_ruolo = [
            partecipante(
                indice,
                regione=f"R{indice % 6}",
                ruolo="Categoria rara" if indice < 6 else "Categoria comune",
                sesso="F" if indice % 2 else "M",
            )
            for indice in range(60)
        ]
        _, gruppi_ruolo = crea_suddivisione(dati_ruolo, 6)
        gruppi_con_categoria_rara = sum(
            any(dati_ruolo[indice]["ruolo"] == "Categoria rara" for indice in gruppo)
            for gruppo in gruppi_ruolo
        )
        self.assertEqual(gruppi_con_categoria_rara, 6)

    def test_normalizzazione_foca_e_sesso(self):
        self.assertEqual(normalizza_foca(" Nomina "), "Nomina")
        self.assertEqual(normalizza_foca("cFa"), "CFA")
        self.assertEqual(normalizza_foca(" CFM "), "CFM")
        self.assertEqual(normalizza_foca("qualunque altro valore"), "Sistema")
        self.assertEqual(normalizza_foca(None), "Sistema")
        self.assertEqual(normalizza_sesso(" m "), "M")
        self.assertEqual(normalizza_sesso("F"), "F")
        self.assertEqual(normalizza_sesso(""), "Non indicato")
        self.assertEqual(normalizza_sesso("X"), "Non indicato")

    def test_priorita_lessicografica_regione_prima_degli_altri_criteri(self):
        dati = [
            partecipante(0, regione="R1", ruolo="A", foca="Nomina", sesso="M"),
            partecipante(1, regione="R1", ruolo="B", foca="CFA", sesso="F"),
            partecipante(2, regione="R2", ruolo="A", foca="Nomina", sesso="M"),
            partecipante(3, regione="R2", ruolo="B", foca="CFA", sesso="F"),
        ]
        regione_bilanciata = [[0, 2], [1, 3]]
        criteri_secondari_bilanciati = [[0, 1], [2, 3]]
        self.assertLess(
            punteggio_suddivisione(regione_bilanciata, dati),
            punteggio_suddivisione(criteri_secondari_bilanciati, dati),
        )
        _, gruppi_generati = crea_suddivisione(dati, 2)
        self.assertEqual(punteggio_suddivisione(gruppi_generati, dati)[0], 0)

    def test_riproducibilita(self):
        dati = [
            partecipante(
                indice,
                regione=f"R{indice % 7}",
                ruolo=f"Q{indice % 5}",
                foca=["Nomina", "CFA", "CFM", "Sistema"][indice % 4],
                sesso=["M", "F", "Non indicato"][indice % 3],
            )
            for indice in range(73)
        ]
        prima, _ = crea_suddivisione(dati, 8)
        seconda, _ = crea_suddivisione(dati, 8)
        self.assertEqual(prima, seconda)

    def test_excel_usa_regione_colonna_i_e_produce_output_leggibile(self):
        workbook = crea_workbook_excel(13)
        foglio, riga_intestazioni, dati = estrai_partecipanti_suddivisione(workbook)

        self.assertEqual(len(dati), 13)
        self.assertEqual(riga_intestazioni, 2)
        self.assertEqual(dati[0]["regione"], "Regione 0")
        self.assertNotEqual(dati[0]["regione"], "Regione sbagliata")

        gruppi = genera_excel_suddivisione(workbook, dati, 4)
        contenuto = workbook_bytes(workbook)
        riletto = load_workbook(contenuto, read_only=True)

        self.assertEqual(
            riletto.sheetnames,
            ["Suddivisione", "Riepilogo", "Gruppo 1", "Gruppo 2", "Gruppo 3", "Gruppo 4"],
        )
        primo_foglio = riletto.worksheets[0]
        righe_generali = list(primo_foglio.iter_rows(values_only=True))
        self.assertEqual(
            righe_generali[0],
            (
                "Codice censimento",
                "Nome",
                "Cognome",
                "Email",
                "Regione",
                "Partecipo in qualità di",
                "FoCa",
                "Sesso",
                "Gruppo",
            ),
        )
        self.assertEqual(primo_foglio.max_column, 9)
        self.assertNotIn("Altro 1", righe_generali[0])
        self.assertNotIn("Altro 2", righe_generali[0])
        self.assertEqual([riga[0] for riga in righe_generali[1:]], list(range(1000, 1013)))
        self.assertEqual(
            [riga[3] for riga in righe_generali[1:]],
            [f"persona{indice}@example.test" for indice in range(13)],
        )
        self.assertEqual(righe_generali[1][4], "Regione 0")
        self.assertNotEqual(righe_generali[1][4], "Regione sbagliata")
        self.assertEqual(righe_generali[1][5], "Categoria rara")
        self.assertEqual(righe_generali[1][6], "Nomina")
        self.assertEqual(righe_generali[2][6], "CFA")
        self.assertEqual(righe_generali[3][6], "CFM")
        self.assertEqual(righe_generali[4][6], "Sistema")
        self.assertEqual(righe_generali[1][7], "M")
        self.assertEqual(righe_generali[2][7], "F")
        self.assertEqual(righe_generali[3][7], "Non indicato")
        self.assertEqual(righe_generali[4][7], "Non indicato")
        self.assertTrue(all(isinstance(riga[8], int) for riga in righe_generali[1:]))

        iscritti_generali = {
            riga[0]: (riga[1], riga[2], riga[3], riga[8])
            for riga in righe_generali[1:]
        }
        iscritti_nei_gruppi = []
        for numero_gruppo in range(1, 5):
            righe_gruppo = list(
                riletto[f"Gruppo {numero_gruppo}"].iter_rows(values_only=True)
            )
            self.assertEqual(
                righe_gruppo[0],
                ("Codice censimento", "Nome", "Cognome", "Email"),
            )
            for codice, nome, cognome, email in righe_gruppo[1:]:
                self.assertEqual(
                    (nome, cognome, email, numero_gruppo),
                    iscritti_generali[codice],
                )
                iscritti_nei_gruppi.append(codice)

        self.assertEqual(len(iscritti_nei_gruppi), 13)
        self.assertEqual(len(set(iscritti_nei_gruppi)), 13)
        self.assertEqual(set(iscritti_nei_gruppi), set(iscritti_generali))
        self.assertEqual(sorted(len(gruppo) for gruppo in gruppi), [3, 3, 3, 4])
        righe_riepilogo = list(
            riletto["Riepilogo"].iter_rows(values_only=True)
        )
        intestazioni_riepilogo = righe_riepilogo[0]
        self.assertEqual(intestazioni_riepilogo[0], "Gruppo")
        self.assertEqual(
            [riga[0] for riga in righe_riepilogo[1:]],
            ["Gruppo 1", "Gruppo 2", "Gruppo 3", "Gruppo 4"],
        )
        self.assertIn("Totale partecipanti", intestazioni_riepilogo)
        self.assertIn("Sesso — Non indicato", intestazioni_riepilogo)
        self.assertIn("FoCa — Sistema", intestazioni_riepilogo)
        self.assertIn("Regione — Regione 0", intestazioni_riepilogo)
        self.assertIn("Partecipo in qualità di — Categoria rara", intestazioni_riepilogo)

    def test_codice_censimento_deve_essere_intero_positivo_nel_limite_database(self):
        for valore in ("non numerico", 0, -1, 2147483648):
            with self.subTest(valore=valore):
                workbook = crea_workbook_excel(1)
                workbook.active.cell(3, 10, valore)

                with self.assertRaisesRegex(
                    ErroreSuddivisione,
                    r"Codice censimento non valido.*intero positivo.*2147483647",
                ):
                    estrai_partecipanti_suddivisione(workbook)

    def test_codici_censimento_duplicati_sono_bloccati_ed_elencati(self):
        workbook = crea_workbook_excel(4)
        for numero_riga, codice in enumerate((123, 456, 123, 456), start=3):
            workbook.active.cell(numero_riga, 10, codice)

        with self.assertRaises(ErroreSuddivisione) as contesto:
            estrai_partecipanti_suddivisione(workbook)

        self.assertEqual(
            str(contesto.exception),
            "Codici censimento duplicati nel file: 123, 456.",
        )


class EndpointSuddivisioneTestCase(unittest.TestCase):
    def setUp(self):
        app.config.update(TESTING=True)
        self.app_context = app.app_context()
        self.app_context.push()
        db.create_all()
        db.session.add(User(username="admin", password=generate_password_hash("password")))
        db.session.commit()
        self.client = app.test_client()

    def tearDown(self):
        db.session.remove()
        db.drop_all()
        self.app_context.pop()

    def login_admin(self):
        self.client.post("/login", data={"username": "admin", "passwd": "password"})

    def file_form(self, numero_partecipanti=12):
        return (
            workbook_bytes(crea_workbook_excel(numero_partecipanti)),
            "iscritti.xlsx",
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )

    def test_endpoint_sono_protetti(self):
        self.assertEqual(self.client.get("/admin/suddivisione-gruppi").status_code, 302)
        self.assertEqual(
            self.client.post("/admin/suddivisione-gruppi/valida").status_code,
            302,
        )
        self.assertEqual(
            self.client.post("/admin/suddivisione-gruppi/genera").status_code,
            302,
        )

        db.session.add(
            User(username="operatore", password=generate_password_hash("password"))
        )
        db.session.commit()
        self.client.post(
            "/login",
            data={"username": "operatore", "passwd": "password"},
        )
        self.assertEqual(
            self.client.get("/admin/suddivisione-gruppi").headers["Location"],
            "/",
        )
        self.assertEqual(
            self.client.post("/admin/suddivisione-gruppi/valida").status_code,
            403,
        )
        self.assertEqual(
            self.client.post("/admin/suddivisione-gruppi/genera").status_code,
            403,
        )

    def test_flusso_validazione_e_generazione(self):
        self.login_admin()
        pagina = self.client.get("/admin/suddivisione-gruppi")
        self.assertEqual(pagina.status_code, 200)
        self.assertIn(b"Suddivisione gruppi", pagina.data)
        self.assertIn(b"Ricalcola gruppi domenica", pagina.data)
        self.assertNotIn(b'type="file"', pagina.data)
        self.assertNotIn(b'numero_laboratori', pagina.data)
        self.assertIn(b'href="/admin/iscrizioni"', pagina.data)
        self.assertIn(b'aria-current="page"', pagina.data)

        validazione = self.client.post(
            "/admin/suddivisione-gruppi/valida",
            data={"file": self.file_form(13)},
        )
        self.assertEqual(validazione.status_code, 200)
        self.assertEqual(validazione.get_json()["partecipanti"], 13)

        risposta = self.client.post(
            "/admin/suddivisione-gruppi/genera",
            data={"file": self.file_form(13), "numero_laboratori": "4"},
        )
        self.assertEqual(risposta.status_code, 200)
        self.assertEqual(
            risposta.content_type,
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        self.assertTrue(risposta.data.startswith(b"PK"))
        self.assertEqual(risposta.headers["X-Partecipanti"], "13")
        self.assertEqual(risposta.headers["X-Laboratori"], "4")
        self.assertEqual(risposta.headers["X-Dimensioni-Gruppi"], '{"3":3,"4":1}')
        riletto = load_workbook(io.BytesIO(risposta.data), read_only=True)
        self.assertIn("Riepilogo", riletto.sheetnames)

    def test_intestazioni_reali_sorgente_sono_validate_e_generate(self):
        workbook = crea_workbook_excel_intestazioni_reali()
        _, riga_intestazioni, partecipanti = estrai_partecipanti_suddivisione(workbook)

        self.assertEqual(riga_intestazioni, 1)
        self.assertEqual(len(partecipanti), 1)
        self.assertEqual(
            partecipanti[0],
            {
                "indice": 0,
                "riga_excel": 2,
                "codice": 123,
                "nome": "Mario",
                "cognome": "Rossi",
                "email": "mario@example.test",
                "regione": "Lazio",
                "ruolo": "Capo",
                "foca": "CFA",
                "sesso": "M",
            },
        )
        self.assertNotEqual(partecipanti[0]["regione"], "Regione da ignorare")

        self.login_admin()
        validazione = self.client.post(
            "/admin/suddivisione-gruppi/valida",
            data={
                "file": (
                    workbook_bytes(crea_workbook_excel_intestazioni_reali()),
                    "iscritti_reali.xlsx",
                )
            },
        )
        self.assertEqual(validazione.status_code, 200)
        self.assertEqual(validazione.get_json()["partecipanti"], 1)

        risposta = self.client.post(
            "/admin/suddivisione-gruppi/genera",
            data={
                "file": (
                    workbook_bytes(crea_workbook_excel_intestazioni_reali()),
                    "iscritti_reali.xlsx",
                ),
                "numero_laboratori": "1",
            },
        )
        self.assertEqual(risposta.status_code, 200)
        riletto = load_workbook(io.BytesIO(risposta.data), read_only=True)
        self.assertEqual(
            riletto.sheetnames,
            ["Suddivisione", "Riepilogo", "Gruppo 1"],
        )
        righe_generali = list(
            riletto["Suddivisione"].iter_rows(values_only=True)
        )
        self.assertEqual(
            righe_generali,
            [
                (
                    "Codice censimento",
                    "Nome",
                    "Cognome",
                    "Email",
                    "Regione",
                    "Partecipo in qualità di",
                    "FoCa",
                    "Sesso",
                    "Gruppo",
                ),
                (
                    123,
                    "Mario",
                    "Rossi",
                    "mario@example.test",
                    "Lazio",
                    "Capo",
                    "CFA",
                    "M",
                    1,
                ),
            ],
        )
        righe_gruppo = list(riletto["Gruppo 1"].iter_rows(values_only=True))
        self.assertEqual(
            righe_gruppo,
            [
                ("Codice censimento", "Nome", "Cognome", "Email"),
                (123, "Mario", "Rossi", "mario@example.test"),
            ],
        )

    def test_generazione_non_modifica_il_dominio_iscrizioni(self):
        self.login_admin()
        stato_prima = {
            modello.__tablename__: modello.query.count()
            for modello in (User, Partecipante, Laboratorio, Iscrizione, SysOption)
        }

        risposta = self.client.post(
            "/admin/suddivisione-gruppi/genera",
            data={"file": self.file_form(13), "numero_laboratori": "4"},
        )

        self.assertEqual(risposta.status_code, 200)
        stato_dopo = {
            modello.__tablename__: modello.query.count()
            for modello in (User, Partecipante, Laboratorio, Iscrizione, SysOption)
        }
        self.assertEqual(stato_dopo, stato_prima)

    def test_file_e_numero_gruppi_non_validi_sono_bloccati(self):
        self.login_admin()
        for nome, contenuto in (
            ("iscritti.txt", io.BytesIO(b"non excel")),
            ("iscritti.xlsx", io.BytesIO(b"non excel")),
        ):
            with self.subTest(nome=nome):
                risposta = self.client.post(
                    "/admin/suddivisione-gruppi/valida",
                    data={"file": (contenuto, nome)},
                )
                self.assertEqual(risposta.status_code, 400)
                self.assertFalse(risposta.get_json()["ok"])

        senza_partecipanti = crea_workbook_excel(0)
        risposta = self.client.post(
            "/admin/suddivisione-gruppi/valida",
            data={
                "file": (
                    workbook_bytes(senza_partecipanti),
                    "senza_partecipanti.xlsx",
                )
            },
        )
        self.assertEqual(risposta.status_code, 400)
        self.assertIn("non contiene partecipanti", risposta.get_json()["errore"])

        regione_nella_colonna_sbagliata = crea_workbook_excel(3)
        regione_nella_colonna_sbagliata.active.cell(2, 9, "Provincia")
        risposta = self.client.post(
            "/admin/suddivisione-gruppi/valida",
            data={
                "file": (
                    workbook_bytes(regione_nella_colonna_sbagliata),
                    "regione_sbagliata.xlsx",
                )
            },
        )
        self.assertEqual(risposta.status_code, 400)
        self.assertIn("colonna I", risposta.get_json()["errore"])

        for numero in ("0", "-1", "14", "1.5", "abc"):
            with self.subTest(numero=numero):
                risposta = self.client.post(
                    "/admin/suddivisione-gruppi/genera",
                    data={"file": self.file_form(13), "numero_laboratori": numero},
                )
                self.assertEqual(risposta.status_code, 400)
                self.assertFalse(risposta.get_json()["ok"])

    def test_upload_oltre_limite_http_restituisce_errore_leggibile(self):
        self.login_admin()
        self.assertEqual(app.config["MAX_CONTENT_LENGTH"], 10 * 1024 * 1024)

        with io.BytesIO(
            b"x" * (app.config["MAX_CONTENT_LENGTH"] + 1)
        ) as contenuto:
            risposta = self.client.post(
                "/admin/suddivisione-gruppi/valida",
                data={"file": (contenuto, "troppo_grande.xlsx")},
            )

        self.assertEqual(risposta.status_code, 413)
        esito = risposta.get_json()
        self.assertFalse(esito["ok"])
        self.assertIn("10 MB", esito["errore"])

    def test_errori_interni_non_sono_mascherati_come_input_non_valido(self):
        self.login_admin()

        for percorso in (
            "/admin/suddivisione-gruppi/valida",
            "/admin/suddivisione-gruppi/genera",
        ):
            with self.subTest(percorso=percorso), patch(
                "app.carica_excel_suddivisione",
                side_effect=RuntimeError("errore interno simulato"),
            ):
                risposta = self.client.post(percorso)
                self.assertEqual(risposta.status_code, 500)
                esito = risposta.get_json()
                self.assertFalse(esito["ok"])
                self.assertNotIn("simulato", esito["errore"])


if __name__ == "__main__":
    unittest.main()
